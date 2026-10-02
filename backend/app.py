import os
import json
import base64
import sqlite3
import threading
import time
import io
import csv
from datetime import datetime, timedelta
from flask import Flask, request, jsonify, send_from_directory, send_file, g
from werkzeug.security import generate_password_hash, check_password_hash
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    _HAS_OPENPYXL = False
    print("[Excel WARNING] 'openpyxl' no está instalado. Ejecute 'pip install openpyxl' para soporte nativo de hojas de cálculo .xlsx.")

from database import get_db, init_db, DB_PATH, get_peru_now, get_peru_now_str, get_peru_today_str
from crypto_utils import encrypt_vault_secret, decrypt_vault_secret

app = Flask(__name__)

@app.teardown_appcontext
def close_open_connections(exception=None):
    conns = getattr(g, '_open_conns', [])
    for conn in conns:
        try:
            conn.close()
        except Exception:
            pass

UPLOAD_FOLDER = os.environ.get('UPLOAD_FOLDER') or os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


# Helper para normalizar evidencias y tareas compartidas en el diccionario de una actividad
def format_actividad_dict(act_row, users_map=None):
    a_dict = dict(act_row)
    ev_val = a_dict.get('evidencias')
    if isinstance(ev_val, str):
        try:
            a_dict['evidencias'] = json.loads(ev_val)
        except Exception:
            a_dict['evidencias'] = []
    elif isinstance(ev_val, list):
        a_dict['evidencias'] = ev_val
    else:
        a_dict['evidencias'] = []

    sw_val = a_dict.get('shared_with')
    if isinstance(sw_val, str):
        try:
            a_dict['shared_with'] = json.loads(sw_val)
        except Exception:
            a_dict['shared_with'] = []
    elif isinstance(sw_val, list):
        a_dict['shared_with'] = sw_val
    else:
        a_dict['shared_with'] = []

    if users_map and isinstance(a_dict.get('shared_with'), list):
        a_dict['shared_with_names'] = [
            users_map[uid] for uid in a_dict['shared_with'] if uid in users_map
        ]
    return a_dict

# Configuración de CORS manual y limpia
@app.after_request
def add_cors_headers(response):
    response.headers['Access-Control-Allow-Origin'] = '*'
    response.headers['Access-Control-Allow-Headers'] = 'Content-Type,Authorization'
    response.headers['Access-Control-Allow-Methods'] = 'GET,PUT,POST,DELETE,OPTIONS,PATCH'
    return response

def get_default_hora_inicio():
    try:
        conn = get_db()
        row = conn.execute("SELECT value FROM system_settings WHERE key = 'hora_inicio_default'").fetchone()
        conn.close()
        if row and row['value']:
            return row['value']
    except Exception:
        pass
    return '08:30'

@app.route('/api/health', methods=['GET'])
def health_check():
    now_dt = get_peru_now()
    return jsonify({
        "status": "ok",
        "service": "Bitácora API",
        "version": "2.2.0",
        "peru_time": now_dt.strftime('%Y-%m-%d %H:%M:%S'),
        "peru_date": get_peru_today_str(),
        "timezone": "America/Lima (UTC-5)",
        "env_tz": os.environ.get('TZ', 'not_set')
    })


# ==================== CONFIGURACIONES GLOBALES (SETTINGS) ====================

@app.route('/api/settings', methods=['GET', 'POST', 'OPTIONS'])
def handle_settings():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'POST':
        data = request.get_json() or {}
        allowed_keys = [
            'hora_inicio_default', 'system_title', 'logo_url', 'favicon_url',
            'login_bg_url', 'login_bg_type', 'login_heading', 'login_subheading',
            'system_mode', 'show_demo_logins'
        ]
        for key in allowed_keys:
            if key in data:
                cursor.execute('''
                    INSERT INTO system_settings (key, value, updated_at)
                    VALUES (?, ?, CURRENT_TIMESTAMP)
                    ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = CURRENT_TIMESTAMP
                ''', (key, str(data[key]) if data[key] is not None else ''))
        conn.commit()

    rows = cursor.execute("SELECT key, value FROM system_settings").fetchall()
    settings = {r['key']: r['value'] for r in rows}
    
    defaults = {
        'hora_inicio_default': '08:30',
        'system_title': 'Bitácora Diaria de Actividades | Marketing Alterno Perú',
        'logo_url': '',
        'favicon_url': '',
        'login_bg_url': '',
        'login_bg_type': 'gradient',
        'login_heading': 'Bitácora Oficial',
        'login_subheading': 'Marketing Alterno Perú',
        'system_mode': 'production',
        'show_demo_logins': 'false'
    }
    for k, v in defaults.items():
        if k not in settings:
            settings[k] = v

    conn.close()
    return jsonify(settings), 200

@app.route('/api/admin/diagnose-db', methods=['GET'])
def diagnose_db():
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        all_acts = [dict(r) for r in cursor.execute("""
            SELECT a.id, a.bitacora_id, a.descripcion, a.estado, a.is_deleted, a.deleted_at, a.created_at,
                   b.colaborador, b.fecha, b.user_id
            FROM actividades a
            LEFT JOIN bitacoras b ON a.bitacora_id = b.id
            ORDER BY a.id DESC
        """).fetchall()]
        
        all_b = [dict(r) for r in cursor.execute("SELECT * FROM bitacoras ORDER BY id DESC").fetchall()]
        tables = [r[0] for r in cursor.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()]
        
        db_dir = os.path.dirname(os.path.abspath(DB_PATH))
        db_files = []
        try:
            db_files = [f for f in os.listdir(db_dir) if any(ext in f for ext in ['db', 'sqlite', 'wal', 'bak'])]
        except Exception:
            pass
            
        import re
        raw_matches = []
        for fname in ['bitacora.db', 'bitacora.db-wal']:
            fpath = os.path.join(db_dir, fname)
            if os.path.exists(fpath):
                try:
                    with open(fpath, 'rb') as fp:
                        raw_bytes = fp.read()
                    raw_strs = re.findall(b'[\x20-\x7E\xC0-\xFF]{6,}', raw_bytes)
                    for s in raw_strs:
                        try:
                            txt = s.decode('latin1')
                            if any(k in txt.lower() for k in ['masivo', 'rh', 'chamba', 'ayala', 'ticket', 'wifi', 'switch', 'vpn', 'incidente', 'odoo']):
                                raw_matches.append(f"[{fname}] {txt}")
                        except Exception:
                            pass
                except Exception as e:
                    raw_matches.append(f"Error reading {fname}: {str(e)}")

        conn.close()
        return jsonify({
            "db_path": DB_PATH,
            "db_files": db_files,
            "tables": tables,
            "total_actividades": len(all_acts),
            "actividades": all_acts,
            "bitacoras": all_b,
            "raw_matches_count": len(raw_matches),
            "raw_matches": list(set(raw_matches))[:200]
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

def cleanup_synthetic_historical_data():
    """
    Elimina bitácoras y actividades sintéticas/falsas inyectadas por scripts antiguos
    para la fecha 2026-09-23 creadas a las 19:40:16 UTC que confunden al usuario administrador.
    """
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        # Eliminar actividades asociadas a bitácoras sintéticas del 2026-09-23 creadas a las 19:40 UTC
        cursor.execute("""
            DELETE FROM actividades 
            WHERE bitacora_id IN (
                SELECT id FROM bitacoras 
                WHERE fecha = '2026-09-23' AND created_at LIKE '2026-09-24 19:40:%'
            )
        """)
        
        # Eliminar las bitácoras sintéticas del 2026-09-23
        cursor.execute("""
            DELETE FROM bitacoras 
            WHERE fecha = '2026-09-23' AND created_at LIKE '2026-09-24 19:40:%'
        """)

        
        conn.commit()
        conn.close()
        print("[Cleanup] Bitácoras sintéticas del 2026-09-23 purgadas exitosamente.")
    except Exception as e:
        print("[Cleanup] Error en cleanup_synthetic_historical_data:", e)

@app.route('/api/admin/clean-synthetic-history', methods=['POST', 'GET', 'OPTIONS'])
def clean_synthetic_history_endpoint():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    cleanup_synthetic_historical_data()
    return jsonify({"success": True, "message": "Bitácoras sintéticas de prueba del 2026-09-23 purgadas correctamente."})

def fix_mock_data_jayala():
    """Limpia las actividades mock generadas por el script de sincronización para Josué Ayala y registra su actividad real 're rh masivo'"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        now_peru = get_peru_now_str()
        
        b_jayala = cursor.execute("""
            SELECT id FROM bitacoras 
            WHERE (user_id = 5 OR LOWER(colaborador) = 'josue ayala') AND fecha = '2026-09-24'
        """).fetchone()

        
        if b_jayala:
            b_id = b_jayala['id']
            # Desactivar actividades mock generadas automáticamente
            mock_descriptions = [
                'Test actividad de depuracion',
                'Implementación de optimizaciones en endpoints API REST y validaciones de carga.',
                'Pruebas integrales de flujo de bitácoras y resolución de inconsistencias de navegación.',
                'Revisión técnica de incidencias y sincronización con el equipo de soporte.',
                'Desarrollo de módulos interactivos de supervisión y gestión de estados de tareas.'
            ]
            for desc in mock_descriptions:
                cursor.execute("""
                    UPDATE actividades SET is_deleted = 1, deleted_at = ?
                    WHERE bitacora_id = ? AND descripcion = ? AND (is_deleted IS NULL OR is_deleted = 0)
                """, (now_peru, b_id, desc))
            
            exact_description = """Se agregó soporte para Carné de Extranjería (CE) como tipo de documento alternativo al DNI en el módulo RH Masivo (mkt_rh_management): nuevo campo "Tipo de Documento" en las líneas de importación, con validaciones diferenciadas (DNI exige 8 dígitos y validación RENIEC; CE exige 9 dígitos con validación manual). Se mantienen intactas las validaciones de correo, celular, provincia, cargo y fechas para ambos tipos.

Durante las pruebas en el sistema se detectaron y corrigieron 2 bugs relacionados:
- Una restricción dura del DNI rompía la carga completa del Excel cuando una fila tenía el documento mal formateado, dejando un mensaje de error pegado en la cabecera de la planilla aunque la fila ya estuviera corregida. Se eliminó, dejando solo la validación que marca la fila puntual en rojo sin bloquear el resto de la carga.
- La plantilla oficial de Excel tenía la columna N° Documento en formato numérico, por lo que Excel borraba los ceros a la izquierda de los Carné de Extranjería antes de guardarse. Se corrigió el formato de esa columna a Texto y se agregó un mensaje de error más claro para detectar este caso a futuro."""

            # Actualizar o insertar con la descripción exacta y fidedigna
            existing_real = cursor.execute("""
                SELECT id FROM actividades 
                WHERE bitacora_id = ? AND (LOWER(descripcion) LIKE '%rh masivo%' OR LOWER(descripcion) LIKE '%carné de extranjería%')
                ORDER BY id DESC LIMIT 1
            """, (b_id,)).fetchone()

            if existing_real:
                cursor.execute("""
                    UPDATE actividades SET 
                        descripcion = ?, duracion_min = 240, para_cliente = 'mkt_rh_management / RR.HH.',
                        tipo_trabajo = 'Desarrollo', estado = 'completada', is_deleted = 0, deleted_at = NULL,
                        updated_at = ?
                    WHERE id = ?
                """, (exact_description, now_peru, existing_real['id']))
            else:
                cursor.execute("""
                    INSERT INTO actividades (
                        bitacora_id, orden, hora_inicio, duracion_min,
                        tipo_trabajo, descripcion, para_cliente, estado, evidencias,
                        shared_with, created_at, updated_at
                    ) VALUES (
                        ?, 0, '08:30', 240,
                        'Desarrollo', ?, 'mkt_rh_management / RR.HH.', 'completada', '[]',
                        '[]', ?, ?
                    )
                """, (b_id, exact_description, now_peru, now_peru))
            
            cursor.execute("""
                UPDATE bitacoras SET tiempo_total_min = (
                    SELECT COALESCE(SUM(duracion_min), 0) FROM actividades WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)
                ), updated_at = ? WHERE id = ?
            """, (b_id, now_peru, b_id))

            
            conn.commit()
        conn.close()
    except Exception as e:
        print("[Repair] Error in fix_mock_data_jayala:", e)

@app.route('/api/admin/repair-mock-jayala', methods=['POST', 'GET', 'OPTIONS'])
def repair_mock_jayala_endpoint():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    fix_mock_data_jayala()
    conn = get_db()
    cursor = conn.cursor()
    acts = cursor.execute("""
        SELECT a.* FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        WHERE (b.user_id = 5 OR LOWER(b.colaborador) = 'josue ayala') AND b.fecha = '2026-09-24' AND (a.is_deleted IS NULL OR a.is_deleted = 0)
    """).fetchall()
    conn.close()
    return jsonify({
        "success": True,
        "message": "Datos mock de Josué Ayala limpiados y actividad real 're rh masivo' registrada con éxito",
        "actividades": [dict(r) for r in acts]
    }), 200

@app.route('/api/admin/clean-production-data', methods=['POST', 'OPTIONS'])
def clean_production_data():
    """Purga todas las bitácoras, actividades y usuarios de prueba dejando el sistema limpio en producción con sólo el admin"""
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute("UPDATE actividades SET parent_task_id = NULL")
        cursor.execute("UPDATE users SET team_id = NULL")
        cursor.execute("UPDATE teams SET lider_id = NULL")
        
        cursor.execute("DELETE FROM actividades")
        cursor.execute("DELETE FROM bitacoras")
        cursor.execute("DELETE FROM buzon_votos")
        cursor.execute("DELETE FROM buzon_sugerencias")
        
        for tbl in ['tasks', 'audit_logs', 'notifications', 'evidencias']:
            try:
                cursor.execute(f"DELETE FROM {tbl}")
            except Exception:
                pass
                
        cursor.execute("DELETE FROM teams")
        cursor.execute("DELETE FROM users WHERE username != 'admin'")
        
        cursor.execute('''
            INSERT INTO system_settings (key, value, updated_at) VALUES ('system_mode', 'production', CURRENT_TIMESTAMP)
            ON CONFLICT(key) DO UPDATE SET value = 'production', updated_at = CURRENT_TIMESTAMP
        ''')
        cursor.execute('''
            INSERT INTO system_settings (key, value, updated_at) VALUES ('show_demo_logins', 'false', CURRENT_TIMESTAMP)
            ON CONFLICT(key) DO UPDATE SET value = 'false', updated_at = CURRENT_TIMESTAMP
        ''')
        conn.commit()
        conn.close()
        return jsonify({
            "success": True,
            "message": "Datos de prueba purgados exitosamente. El sistema está limpio en Modo Producción con el usuario Administrador."
        }), 200
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({"error": f"Error al limpiar datos: {str(e)}"}), 500

@app.route('/api/admin/seed-demo-data', methods=['POST', 'OPTIONS'])
def seed_demo_data():
    """Restaura los datos de prueba (usuarios, equipos y bitácoras de ejemplo) para modo demostración"""
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    try:
        from seed_data import seed_team_and_activities
        seed_team_and_activities()
        
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO system_settings (key, value, updated_at) VALUES ('system_mode', 'demo', CURRENT_TIMESTAMP)
            ON CONFLICT(key) DO UPDATE SET value = 'demo', updated_at = CURRENT_TIMESTAMP
        ''')
        cursor.execute('''
            INSERT INTO system_settings (key, value, updated_at) VALUES ('show_demo_logins', 'true', CURRENT_TIMESTAMP)
            ON CONFLICT(key) DO UPDATE SET value = 'true', updated_at = CURRENT_TIMESTAMP
        ''')
        conn.commit()
        conn.close()
        return jsonify({
            "success": True,
            "message": "Datos de demostración y usuarios de prueba cargados correctamente en Modo Demo."
        }), 200
    except Exception as e:
        return jsonify({"error": f"Error al cargar datos de demostración: {str(e)}"}), 500


# ==================== GESTIÓN DE ARCHIVOS Y EVIDENCIAS ====================

@app.route('/api/upload', methods=['POST', 'OPTIONS'])
def upload_file():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    try:
        # Caso 1: Archivo enviado vía multipart/form-data (explorador o drag & drop)
        if 'file' in request.files:
            file = request.files['file']
            if not file or file.filename == '':
                return jsonify({"error": "No se seleccionó ningún archivo"}), 400
                
            orig_filename = secure_filename(file.filename) or 'archivo'
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            unique_name = f"{timestamp}_{orig_filename}"
            filepath = os.path.join(UPLOAD_FOLDER, unique_name)
            file.save(filepath)
            
            file_size = os.path.getsize(filepath)
            mime_type = file.content_type or 'application/octet-stream'
            
            return jsonify({
                "url": f"/api/uploads/{unique_name}",
                "nombre": orig_filename,
                "tipo": mime_type,
                "tamano": file_size
            }), 201

        # Caso 2: Imagen pegada como Base64 (capturas con Ctrl+V)
        data = request.get_json() or {}
        base64_data = data.get('base64')
        if base64_data:
            nombre = data.get('nombre') or 'captura_pantalla.png'
            safe_nombre = secure_filename(nombre) or 'captura.png'
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            unique_name = f"{timestamp}_{safe_nombre}"
            filepath = os.path.join(UPLOAD_FOLDER, unique_name)
            
            if ',' in base64_data:
                header, base64_str = base64_data.split(',', 1)
                mime_type = header.split(';')[0].replace('data:', '') if 'data:' in header else 'image/png'
            else:
                base64_str = base64_data
                mime_type = 'image/png'
                
            binary_bytes = base64.b64decode(base64_str)
            with open(filepath, 'wb') as f:
                f.write(binary_bytes)
                
            file_size = len(binary_bytes)
            
            return jsonify({
                "url": f"/api/uploads/{unique_name}",
                "nombre": safe_nombre,
                "tipo": mime_type,
                "tamano": file_size
            }), 201

        return jsonify({"error": "No se proporcionó archivo ni imagen base64"}), 400
    except Exception as e:
        return jsonify({"error": f"Error al procesar el archivo: {str(e)}"}), 500

@app.route('/api/uploads/<path:filename>', methods=['GET'])
def serve_uploaded_file(filename):
    return send_from_directory(UPLOAD_FOLDER, filename)

# ==================== AUTHENTICATION ====================

@app.route('/api/auth/login', methods=['POST', 'OPTIONS'])
def login():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    data = request.get_json() or {}
    username = (data.get('username') or '').strip().lower()
    password = data.get('password') or ''
    
    if not username or not password:
        return jsonify({"error": "Usuario y contraseña requeridos"}), 400
        
    conn = get_db()
    cursor = conn.cursor()
    user = cursor.execute('''
        SELECT u.id, u.username, u.email, u.password_hash, u.full_name, u.role, u.is_active, u.team_id,
               t.nombre as team_name, t.lider_id
        FROM users u
        LEFT JOIN teams t ON u.team_id = t.id
        WHERE LOWER(u.username) = ?
    ''', (username,)).fetchone()
    conn.close()
    
    if not user:
        return jsonify({"error": "Credenciales inválidas"}), 401
        
    if not user['is_active']:
        return jsonify({"error": "Usuario desactivado. Contacta al administrador"}), 403
        
    stored_hash = user['password_hash']
    # Comprobar contraseña con Werkzeug o fallback para desarrollo
    password_matches = False
    try:
        password_matches = check_password_hash(stored_hash, password)
    except Exception:
        password_matches = (stored_hash == password)
        
    if not password_matches and stored_hash == password:
        password_matches = True
        
    if not password_matches:
        return jsonify({"error": "Contraseña incorrecta"}), 401
        
    role_name = 'analista' if user['role'] in ('operador', 'analista') else user['role']
    is_leader = (user['lider_id'] == user['id']) or (role_name == 'lider') or (role_name == 'admin')
    user_dict = {
        "id": user['id'],
        "username": user['username'],
        "email": user['email'],
        "full_name": user['full_name'],
        "role": role_name,
        "phone": user['phone'] if 'phone' in user.keys() else None,
        "team_id": user['team_id'],
        "team_name": user['team_name'],
        "is_leader": is_leader
    }
    
    return jsonify({
        "message": "Autenticación exitosa",
        "user": user_dict,
        "token": f"session-{user['id']}-{user['username']}"
    }), 200

@app.route('/api/auth/change-password', methods=['POST', 'OPTIONS'])
def change_password():
    """Permite a cualquier usuario cambiar su contraseña validando la anterior, o al admin actualizarla."""
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    data = request.get_json() or {}
    user_id = data.get('user_id')
    current_password = data.get('current_password') or ''
    new_password = (data.get('new_password') or '').strip()
    
    if not user_id:
        return jsonify({"error": "Identificador de usuario requerido"}), 400
        
    if not new_password or len(new_password) < 4:
        return jsonify({"error": "La nueva contraseña debe tener al menos 4 caracteres"}), 400
        
    conn = get_db()
    cursor = conn.cursor()
    user = cursor.execute("SELECT id, username, password_hash FROM users WHERE id = ?", (user_id,)).fetchone()
    
    if not user:
        conn.close()
        return jsonify({"error": "Usuario no encontrado"}), 404
        
    stored_hash = user['password_hash']
    # Si se envía contraseña actual, validarla
    if current_password:
        password_matches = False
        try:
            password_matches = check_password_hash(stored_hash, current_password)
        except Exception:
            password_matches = (stored_hash == current_password)
        if not password_matches and stored_hash == current_password:
            password_matches = True
            
        if not password_matches:
            conn.close()
            return jsonify({"error": "La contraseña actual es incorrecta"}), 400
            
    new_hash = generate_password_hash(new_password)
    cursor.execute("UPDATE users SET password_hash = ? WHERE id = ?", (new_hash, user_id))
    conn.commit()
    conn.close()
    
    return jsonify({
        "success": True,
        "message": "Contraseña actualizada exitosamente"
    }), 200

@app.route('/api/auth/profile', methods=['PUT', 'OPTIONS'])
def update_profile():
    """Permite al propio usuario actualizar su nombre, correo, teléfono y opcionalmente su contraseña sin poder modificar roles ni equipos."""
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    data = request.get_json() or {}
    user_id = data.get('user_id')
    if not user_id:
        return jsonify({"error": "Identificador de usuario requerido"}), 400
        
    full_name = (data.get('full_name') or '').strip()
    email = (data.get('email') or '').strip()
    phone = (data.get('phone') or '').strip()
    current_password = data.get('current_password') or ''
    new_password = (data.get('new_password') or '').strip()
    
    if not full_name:
        return jsonify({"error": "El nombre completo es obligatorio"}), 400
        
    conn = get_db()
    cursor = conn.cursor()
    u = cursor.execute("SELECT id, username, password_hash, role, team_id, is_active FROM users WHERE id = ?", (user_id,)).fetchone()
    if not u:
        conn.close()
        return jsonify({"error": "Usuario no encontrado"}), 404
        
    # Si desea cambiar contraseña
    if new_password:
        if len(new_password) < 4:
            conn.close()
            return jsonify({"error": "La nueva contraseña debe tener al menos 4 caracteres"}), 400
        if not current_password:
            conn.close()
            return jsonify({"error": "Debes ingresar tu contraseña actual para establecer una nueva"}), 400
            
        stored_hash = u['password_hash']
        password_matches = False
        try:
            password_matches = check_password_hash(stored_hash, current_password)
        except Exception:
            password_matches = (stored_hash == current_password)
        if not password_matches and stored_hash == current_password:
            password_matches = True
            
        if not password_matches:
            conn.close()
            return jsonify({"error": "La contraseña actual es incorrecta"}), 400
            
        new_hash = generate_password_hash(new_password)
        cursor.execute("""
            UPDATE users 
            SET full_name = ?, email = ?, phone = ?, password_hash = ?
            WHERE id = ?
        """, (full_name, email, phone, new_hash, user_id))
    else:
        cursor.execute("""
            UPDATE users 
            SET full_name = ?, email = ?, phone = ?
            WHERE id = ?
        """, (full_name, email, phone, user_id))
        
    conn.commit()
    
    updated_u = cursor.execute("""
        SELECT u.id, u.username, u.email, u.full_name, u.role, u.phone, u.team_id, u.is_active,
               t.nombre as team_name,
               CASE WHEN t.lider_id = u.id THEN 1 ELSE 0 END as is_team_leader
        FROM users u
        LEFT JOIN teams t ON u.team_id = t.id
        WHERE u.id = ?
    """, (user_id,)).fetchone()
    
    conn.close()
    
    return jsonify({
        "success": True,
        "message": "Perfil actualizado correctamente",
        "user": dict(updated_u)
    }), 200

# ==================== USERS CRUD ====================

@app.route('/api/users', methods=['GET'])
def get_users_list():
    conn = get_db()
    users = conn.execute('''
        SELECT u.id, u.username, u.email, u.full_name, u.role, u.phone, u.team_id, u.is_active,
               t.nombre as team_name,
               CASE WHEN t.lider_id = u.id THEN 1 ELSE 0 END as is_team_leader
        FROM users u
        LEFT JOIN teams t ON u.team_id = t.id
        WHERE u.is_active = 1
        ORDER BY u.full_name ASC
    ''').fetchall()
    conn.close()
    return jsonify([dict(u) for u in users])

@app.route('/api/admin/users', methods=['GET', 'POST', 'OPTIONS'])
def admin_users():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    
    if request.method == 'GET':
        users = cursor.execute('''
            SELECT u.id, u.username, u.email, u.full_name, u.role, u.phone, u.team_id, u.is_active, u.created_at,
                   t.nombre as team_name,
                   CASE WHEN t.lider_id = u.id THEN 1 ELSE 0 END as is_team_leader
            FROM users u
            LEFT JOIN teams t ON u.team_id = t.id
            ORDER BY u.id ASC
        ''').fetchall()
        conn.close()
        return jsonify([dict(u) for u in users])
        
    elif request.method == 'POST':
        data = request.get_json() or {}
        username = (data.get('username') or '').strip().lower()
        full_name = (data.get('full_name') or '').strip()
        email = (data.get('email') or '').strip()
        phone = (data.get('phone') or '').strip()
        password = data.get('password') or '123456'
        role = data.get('role') or 'analista'
        team_id = data.get('team_id') or None
        
        if not username or not full_name:
            conn.close()
            return jsonify({"error": "Usuario y Nombre completo son requeridos"}), 400
            
        existing = cursor.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()
        if existing:
            conn.close()
            return jsonify({"error": f"El nombre de usuario '{username}' ya existe"}), 400
            
        p_hash = generate_password_hash(password)
        cursor.execute('''
            INSERT INTO users (username, email, phone, password_hash, full_name, role, team_id, is_active, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, 1, CURRENT_TIMESTAMP)
        ''', (username, email, phone, p_hash, full_name, role, team_id))
        user_id = cursor.lastrowid
        conn.commit()
        
        new_u = cursor.execute("SELECT id, username, email, phone, full_name, role, team_id FROM users WHERE id = ?", (user_id,)).fetchone()
        conn.close()
        return jsonify({"message": "Usuario creado con éxito", "user": dict(new_u)}), 201

@app.route('/api/admin/users/<int:user_id>', methods=['PUT', 'DELETE', 'OPTIONS'])
def admin_user_detail(user_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    
    if request.method == 'PUT':
        data = request.get_json() or {}
        raw_username = (data.get('username') or '').strip().lower()
        full_name = data.get('full_name')
        email = data.get('email')
        phone = data.get('phone')
        role = data.get('role')
        team_id = data.get('team_id')
        is_active = data.get('is_active', 1)
        password = data.get('password')
        
        current_u = cursor.execute("SELECT id, username FROM users WHERE id = ?", (user_id,)).fetchone()
        if not current_u:
            conn.close()
            return jsonify({"error": "Usuario no encontrado"}), 404
            
        target_username = current_u['username']
        if raw_username and raw_username != current_u['username'].lower():
            if current_u['username'] == 'admin':
                conn.close()
                return jsonify({"error": "No se puede renombrar el usuario de la cuenta root admin"}), 400
            # Validar que no exista otro con ese username
            dup = cursor.execute("SELECT id FROM users WHERE LOWER(username) = ? AND id != ?", (raw_username, user_id)).fetchone()
            if dup:
                conn.close()
                return jsonify({"error": f"El nombre de usuario '{raw_username}' ya está en uso por otro colaborador"}), 400
            target_username = raw_username

        if password and len(password.strip()) > 0:
            p_hash = generate_password_hash(password.strip())
            cursor.execute('''
                UPDATE users SET
                    username = ?,
                    full_name = COALESCE(?, full_name),
                    email = COALESCE(?, email),
                    phone = COALESCE(?, phone),
                    role = COALESCE(?, role),
                    team_id = ?,
                    is_active = ?,
                    password_hash = ?
                WHERE id = ?
            ''', (target_username, full_name, email, phone, role, team_id, is_active, p_hash, user_id))
        else:
            cursor.execute('''
                UPDATE users SET
                    username = ?,
                    full_name = COALESCE(?, full_name),
                    email = COALESCE(?, email),
                    phone = COALESCE(?, phone),
                    role = COALESCE(?, role),
                    team_id = ?,
                    is_active = ?
                WHERE id = ?
            ''', (target_username, full_name, email, phone, role, team_id, is_active, user_id))
            
        conn.commit()
        updated = cursor.execute("SELECT id, username, email, phone, full_name, role, team_id, is_active FROM users WHERE id = ?", (user_id,)).fetchone()
        conn.close()
        return jsonify({"message": "Usuario actualizado con éxito", "user": dict(updated)})
        
    elif request.method == 'DELETE':
        if user_id == 1:
            conn.close()
            return jsonify({"error": "No se puede eliminar ni desactivar al Administrador Principal"}), 400
            
        permanent = request.args.get('permanent', 'false').lower() in ('true', '1', 'yes')
        mode = request.args.get('mode', '')
        if mode == 'permanent':
            permanent = True
            
        if permanent:
            try:
                # 1. Desvincular liderazgo en equipos
                cursor.execute("UPDATE teams SET lider_id = NULL WHERE lider_id = ?", (user_id,))
                # 2. Desvincular dependencias en actividades (tanto hacia sus tareas como de sus tareas)
                cursor.execute('''
                    UPDATE actividades SET parent_task_id = NULL 
                    WHERE parent_task_id IN (SELECT id FROM actividades WHERE bitacora_id IN (SELECT id FROM bitacoras WHERE user_id = ?))
                ''', (user_id,))
                cursor.execute('''
                    UPDATE actividades SET parent_task_id = NULL 
                    WHERE bitacora_id IN (SELECT id FROM bitacoras WHERE user_id = ?)
                ''', (user_id,))
                # 3. Eliminar actividades y bitácoras asociadas
                cursor.execute('''
                    DELETE FROM actividades 
                    WHERE bitacora_id IN (SELECT id FROM bitacoras WHERE user_id = ?)
                ''', (user_id,))
                cursor.execute("DELETE FROM bitacoras WHERE user_id = ?", (user_id,))
                
                for tbl in ['tasks', 'audit_logs', 'notifications']:
                    try:
                        cursor.execute(f"DELETE FROM {tbl} WHERE user_id = ?", (user_id,))
                    except Exception:
                        pass
                        
                # 4. Eliminar el usuario definitivamente
                cursor.execute("DELETE FROM users WHERE id = ?", (user_id,))
                conn.commit()
                conn.close()
                return jsonify({"success": True, "message": "Usuario eliminado definitivamente del sistema"})
            except Exception as e:
                conn.rollback()
                conn.close()
                return jsonify({"error": f"Error al eliminar usuario: {str(e)}"}), 500
        else:
            # Alternar o desactivar
            curr = cursor.execute("SELECT is_active FROM users WHERE id = ?", (user_id,)).fetchone()
            if not curr:
                conn.close()
                return jsonify({"error": "Usuario no encontrado"}), 404
            new_active = 0 if curr['is_active'] == 1 else 1
            msg = "Usuario desactivado correctamente" if new_active == 0 else "Usuario reactivado correctamente"
            cursor.execute("UPDATE users SET is_active = ? WHERE id = ?", (new_active, user_id))
            conn.commit()
            conn.close()
            return jsonify({"success": True, "message": msg, "is_active": new_active})

# ==================== TEAMS CRUD ====================

@app.route('/api/admin/teams', methods=['GET', 'POST', 'OPTIONS'])
def admin_teams():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    
    if request.method == 'GET':
        teams = cursor.execute('''
            SELECT t.id, t.nombre, t.descripcion, t.lider_id, t.created_at,
                   u.full_name as lider_nombre, u.username as lider_username
            FROM teams t
            LEFT JOIN users u ON t.lider_id = u.id
            ORDER BY t.nombre ASC
        ''').fetchall()
        
        result = []
        for t in teams:
            t_dict = dict(t)
            members = cursor.execute('''
                SELECT id, username, full_name, role, email, phone
                FROM users
                WHERE team_id = ? AND is_active = 1
            ''', (t['id'],)).fetchall()
            t_dict['members'] = [dict(m) for m in members]
            t_dict['members_count'] = len(members)
            result.append(t_dict)
            
        conn.close()
        return jsonify(result)
        
    elif request.method == 'POST':
        data = request.get_json() or {}
        nombre = (data.get('nombre') or '').strip()
        descripcion = data.get('descripcion', '')
        lider_id = data.get('lider_id') or None
        
        if not nombre:
            conn.close()
            return jsonify({"error": "El nombre del equipo es requerido"}), 400
            
        cursor.execute('''
            INSERT INTO teams (nombre, descripcion, lider_id)
            VALUES (?, ?, ?)
        ''', (nombre, descripcion, lider_id))
        team_id = cursor.lastrowid
        
        # Si se asignó un líder, actualizar su rol a líder si no es admin
        if lider_id:
            cursor.execute('''
                UPDATE users SET team_id = ?, role = CASE WHEN role = 'admin' THEN 'admin' ELSE 'lider' END
                WHERE id = ?
            ''', (team_id, lider_id))
            
        conn.commit()
        conn.close()
        return jsonify({"message": "Equipo creado con éxito", "team_id": team_id}), 201

@app.route('/api/admin/teams/<int:team_id>', methods=['PUT', 'DELETE', 'OPTIONS'])
def admin_team_detail(team_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    
    if request.method == 'PUT':
        data = request.get_json() or {}
        nombre = data.get('nombre')
        descripcion = data.get('descripcion')
        lider_id = data.get('lider_id')
        member_ids = data.get('member_ids') # opcional lista de IDs para asignar
        
        cursor.execute('''
            UPDATE teams SET
                nombre = COALESCE(?, nombre),
                descripcion = COALESCE(?, descripcion),
                lider_id = ?
            WHERE id = ?
        ''', (nombre, descripcion, lider_id, team_id))
        
        if lider_id:
            cursor.execute('''
                UPDATE users SET team_id = ?, role = CASE WHEN role = 'admin' THEN 'admin' ELSE 'lider' END
                WHERE id = ?
            ''', (team_id, lider_id))
            
        if member_ids is not None and isinstance(member_ids, list):
            # Asignar miembros indicados a este equipo
            cursor.execute("UPDATE users SET team_id = NULL WHERE team_id = ?", (team_id,))
            if member_ids:
                placeholders = ','.join(['?'] * len(member_ids))
                cursor.execute(f"UPDATE users SET team_id = ? WHERE id IN ({placeholders})", [team_id] + member_ids)
                
        conn.commit()
        conn.close()
        return jsonify({"message": "Equipo actualizado correctamente"})
        
    elif request.method == 'DELETE':
        cursor.execute("UPDATE users SET team_id = NULL WHERE team_id = ?", (team_id,))
        cursor.execute("DELETE FROM teams WHERE id = ?", (team_id,))
        conn.commit()
        conn.close()
        return jsonify({"message": "Equipo eliminado correctamente"})

# ==================== BITACORAS CRUD & TEAM SUPERVISION ====================

def rollover_user_open_tasks(cursor, u_id, u_name, today_str, now_peru):
    """
    Arrastra tareas abiertas de días previos de un usuario hacia la fecha actual.
    Idempotente: no duplica si ya fue transferida a hoy o continuada en una fecha posterior.
    """
    open_tasks = cursor.execute('''
        SELECT a.*, b.fecha, b.area
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        WHERE (b.user_id = ? OR b.colaborador = ?)
          AND b.fecha < ?
          AND a.estado != 'completada'
          AND (a.is_deleted IS NULL OR a.is_deleted = 0)
          AND NOT EXISTS (
            SELECT 1 FROM actividades a_child
            JOIN bitacoras b_child ON a_child.bitacora_id = b_child.id
            WHERE a_child.parent_task_id = a.id
              AND (a_child.is_deleted IS NULL OR a_child.is_deleted = 0)
              AND b_child.fecha > b.fecha
          )
        ORDER BY a.id ASC
    ''', (u_id, u_name, today_str)).fetchall()
    
    if not open_tasks:
        return 0
        
    # Buscar si ya existe bitácora de hoy para este colaborador
    today_b = cursor.execute('''
        SELECT id FROM bitacoras
        WHERE (user_id = ? OR colaborador = ?) AND fecha = ?
    ''', (u_id, u_name, today_str)).fetchone()
    
    today_b_id = None
    if today_b:
        today_b_id = today_b['id']
    else:
        last_b = cursor.execute('''
            SELECT area FROM bitacoras
            WHERE (user_id = ? OR colaborador = ?)
            ORDER BY fecha DESC, id DESC LIMIT 1
        ''', (u_id, u_name)).fetchone()
        area_val = last_b['area'] if last_b else 'Sistemas'
        
        cursor.execute('''
            INSERT INTO bitacoras (
                user_id, fecha, hora_inicio, colaborador, area,
                pendientes, necesita_apoyo, prioridad_siguiente,
                tiempo_total_min, estado, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, '', 'No', '', 0, 'generada', ?, ?)
        ''', (u_id, today_str, get_default_hora_inicio(), u_name, area_val, now_peru, now_peru))
        today_b_id = cursor.lastrowid

    # IDs y descripciones de tareas que ya existen hoy
    existing_parent_ids = set(
        r[0] for r in cursor.execute(
            "SELECT parent_task_id FROM actividades WHERE bitacora_id = ? AND parent_task_id IS NOT NULL AND (is_deleted IS NULL OR is_deleted = 0)",
            (today_b_id,)
        ).fetchall()
    )
    existing_descriptions = set(
        r[0].strip().lower() for r in cursor.execute(
            "SELECT descripcion FROM actividades WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)",
            (today_b_id,)
        ).fetchall() if r[0]
    )
    
    current_order = cursor.execute(
        "SELECT COALESCE(MAX(orden), 0) FROM actividades WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)",
        (today_b_id,)
    ).fetchone()[0] + 1
    
    rolled_count = 0
    seen_in_batch = set()
    for task in open_tasks:
        desc_clean = (task['descripcion'] or '').strip().lower()
        if not desc_clean:
            continue
        if task['id'] in existing_parent_ids:
            continue
        if desc_clean in existing_descriptions or desc_clean in seen_in_batch:
            continue

        seen_in_batch.add(desc_clean)
        existing_descriptions.add(desc_clean)
        
        orig_created = task['created_at'] or get_peru_now().strftime('%Y-%m-%d %H:%M:%S')
        orig_updated = task['updated_at'] or orig_created
        ev_val = task['evidencias'] if ('evidencias' in task.keys() and task['evidencias']) else '[]'
        sw_val = task['shared_with'] if ('shared_with' in task.keys() and task['shared_with']) else '[]'
        sw_uuid = task['shared_uuid'] if ('shared_uuid' in task.keys() and task['shared_uuid']) else None
        cursor.execute('''
            INSERT INTO actividades (
                bitacora_id, orden, hora_inicio, duracion_min,
                tipo_trabajo, descripcion, para_cliente, estado,
                parent_task_id, tipo_vinculo, created_at, updated_at, evidencias,
                shared_with, shared_uuid, is_deleted
            ) VALUES (?, ?, '', 0, ?, ?, ?, ?, ?, 'continuacion', ?, ?, ?, ?, ?, 0)
        ''', (
            today_b_id,
            current_order,
            task['tipo_trabajo'],
            task['descripcion'],
            task['para_cliente'],
            task['estado'],
            task['id'],
            orig_created,
            orig_updated,
            ev_val,
            sw_val,
            sw_uuid
        ))
        existing_parent_ids.add(task['id'])
        current_order += 1
        rolled_count += 1
            
    return rolled_count

@app.route('/api/bitacoras', methods=['GET'])
def get_bitacoras():
    conn = get_db()
    cursor = conn.cursor()
    
    fecha = request.args.get('fecha')
    colaborador = request.args.get('colaborador')
    team_id = request.args.get('team_id')
    user_id = request.args.get('user_id')
    requesting_user_id = request.args.get('requesting_user_id')
    
    today_str = get_peru_today_str()
    # Si la petición incluye la fecha de hoy o no especifica fecha, asegurar rollover catch-up para el usuario activo
    if not fecha or fecha == today_str:
        target_uid = user_id or requesting_user_id
        if target_uid:
            target_user = cursor.execute("SELECT id, full_name FROM users WHERE id = ?", (target_uid,)).fetchone()
            if target_user:
                rolled = rollover_user_open_tasks(cursor, target_user['id'], target_user['full_name'], today_str, get_peru_now_str())
                if rolled > 0:
                    conn.commit()
        elif colaborador:
            target_user = cursor.execute("SELECT id, full_name FROM users WHERE LOWER(full_name) = LOWER(?) OR LOWER(username) = LOWER(?)", (colaborador, colaborador)).fetchone()
            if target_user:
                rolled = rollover_user_open_tasks(cursor, target_user['id'], target_user['full_name'], today_str, get_peru_now_str())
                if rolled > 0:
                    conn.commit()
    
    query = '''
        SELECT b.*, u.team_id, t.nombre as team_name
        FROM bitacoras b
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE LOWER(full_name) = LOWER(b.colaborador) OR LOWER(username) = LOWER(b.colaborador) LIMIT 1)) = u.id
        LEFT JOIN teams t ON u.team_id = t.id
        WHERE 1=1
    '''
    params = []
    
    # PRIVACIDAD ESTRICTA POR ROL:
    if requesting_user_id:
        req_u = cursor.execute("SELECT id, role, team_id, full_name, username FROM users WHERE id = ?", (requesting_user_id,)).fetchone()
        if req_u:
            r_role = req_u['role']
            if r_role in ('analista', 'operador'):
                # Los analistas/operadores SOLO pueden ver sus propias bitácoras
                query += " AND (b.user_id = ? OR LOWER(b.colaborador) = LOWER(?) OR LOWER(b.colaborador) = LOWER(?))"
                params.extend([req_u['id'], req_u['full_name'], req_u['username']])
            # Si es lider o admin, pueden ver todas las bitácoras o filtrar por team_id si se pasa como parámetro
    
    if fecha:
        query += " AND b.fecha = ?"
        params.append(fecha)
    if colaborador:
        query += " AND (b.colaborador LIKE ? OR u.username LIKE ?)"
        params.extend([f"%{colaborador}%", f"%{colaborador}%"])
    if team_id:
        query += " AND u.team_id = ?"
        params.append(team_id)
    if user_id:
        query += " AND (b.user_id = ? OR u.id = ?)"
        params.extend([user_id, user_id])
    if fecha:
        query += " ORDER BY b.fecha DESC, b.id DESC LIMIT 500"
    else:
        query += " ORDER BY b.fecha DESC, b.id DESC LIMIT 200"
    bitacoras = cursor.execute(query, params).fetchall()
    users_map = {u['id']: u['full_name'] for u in cursor.execute("SELECT id, full_name FROM users").fetchall()}
    
    result = []
    for b in bitacoras:
        b_dict = dict(b)
        acts = cursor.execute('''
            SELECT a.*,
              p.descripcion as parent_task_desc,
              p.estado as parent_task_estado,
              COALESCE((
                WITH RECURSIVE lineage AS (
                  SELECT id, parent_task_id, duracion_min FROM actividades WHERE id = a.id
                  UNION ALL
                  SELECT prev.id, prev.parent_task_id, prev.duracion_min
                  FROM actividades prev
                  JOIN lineage l ON prev.id = l.parent_task_id
                )
                SELECT SUM(duracion_min) FROM lineage
              ), a.duracion_min) AS tiempo_acumulado_min
            FROM actividades a
            LEFT JOIN actividades p ON a.parent_task_id = p.id
            WHERE a.bitacora_id = ? AND (a.is_deleted IS NULL OR a.is_deleted = 0)
            ORDER BY a.orden ASC, a.id ASC
        ''', (b['id'],)).fetchall()
        b_dict['actividades'] = [format_actividad_dict(a, users_map) for a in acts]
        result.append(b_dict)
        
    conn.close()
    return jsonify(result)

@app.route('/api/bitacoras/<int:bitacora_id>', methods=['GET'])
def get_bitacora(bitacora_id):
    conn = get_db()
    cursor = conn.cursor()
    b = cursor.execute("SELECT * FROM bitacoras WHERE id = ?", (bitacora_id,)).fetchone()
    if not b:
        conn.close()
        return jsonify({"error": "Bitácora no encontrada"}), 404
        
    b_dict = dict(b)
    users_map = {u['id']: u['full_name'] for u in cursor.execute("SELECT id, full_name FROM users").fetchall()}
    acts = cursor.execute('''
        SELECT a.*,
          p.descripcion as parent_task_desc,
          p.estado as parent_task_estado,
          COALESCE((
            WITH RECURSIVE lineage AS (
              SELECT id, parent_task_id, duracion_min FROM actividades WHERE id = a.id
              UNION ALL
              SELECT prev.id, prev.parent_task_id, prev.duracion_min
              FROM actividades prev
              JOIN lineage l ON prev.id = l.parent_task_id
            )
            SELECT SUM(duracion_min) FROM lineage
          ), a.duracion_min) AS tiempo_acumulado_min
        FROM actividades a
        LEFT JOIN actividades p ON a.parent_task_id = p.id
        WHERE a.bitacora_id = ? AND (a.is_deleted IS NULL OR a.is_deleted = 0)
        ORDER BY a.orden ASC, a.id ASC
    ''', (bitacora_id,)).fetchall()
    b_dict['actividades'] = [format_actividad_dict(a, users_map) for a in acts]
    conn.close()
    return jsonify(b_dict)

@app.route('/api/bitacoras', methods=['POST', 'OPTIONS'])
def save_bitacora():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    data = request.get_json() or {}
    fecha = data.get('fecha')
    colaborador = data.get('colaborador')
    area = data.get('area', 'Sistemas')
    user_id = data.get('user_id')
    
    if not fecha or not colaborador:
        return jsonify({"error": "Fecha y Colaborador son requeridos"}), 400

    conn = get_db()
    cursor = conn.cursor()
    
    # Canonicalizar user_id y colaborador con la tabla users
    if not user_id:
        u_match = cursor.execute(
            "SELECT id, full_name FROM users WHERE LOWER(full_name) = LOWER(?) OR LOWER(username) = LOWER(?)",
            (colaborador, colaborador)
        ).fetchone()
        if u_match:
            user_id = u_match['id']
            colaborador = u_match['full_name']
    else:
        u_by_id = cursor.execute("SELECT id, full_name FROM users WHERE id = ?", (user_id,)).fetchone()
        if u_by_id:
            colaborador = u_by_id['full_name']

    try:
        now_peru = get_peru_now_str()
        actividades_data = data.get('actividades', [])
        tiempo_total = sum(int(a.get('duracion_min', 0) or 0) for a in actividades_data)
        bitacora_id = data.get('id')
        
        if bitacora_id:
            existing_b = cursor.execute("SELECT id FROM bitacoras WHERE id = ?", (bitacora_id,)).fetchone()
            if existing_b:
                cursor.execute('''
                    UPDATE bitacoras SET
                        user_id = COALESCE(?, user_id),
                        fecha = ?,
                        hora_inicio = ?,
                        colaborador = ?,
                        area = ?,
                        pendientes = ?,
                        necesita_apoyo = ?,
                        apoyo_detalle = ?,
                        prioridad_siguiente = ?,
                        tiempo_total_min = ?,
                        resumen_texto = ?,
                        estado = ?,
                        updated_at = ?
                    WHERE id = ?
                ''', (
                    user_id,
                    fecha,
                    data.get('hora_inicio', ''),
                    colaborador,
                    area,
                    data.get('pendientes', ''),
                    data.get('necesita_apoyo', 'No'),
                    data.get('apoyo_detalle', ''),
                    data.get('prioridad_siguiente', ''),
                    tiempo_total,
                    data.get('resumen_texto', ''),
                    data.get('estado', 'generada'),
                    now_peru,
                    bitacora_id
                ))
            else:
                bitacora_id = None

        if not bitacora_id:
            # Si ya existe una bitácora para este colaborador y fecha, actualizarla para no duplicar
            existing_by_date = cursor.execute(
                "SELECT id FROM bitacoras WHERE (user_id = ? OR LOWER(colaborador) = LOWER(?)) AND fecha = ?",
                (user_id, colaborador, fecha)
            ).fetchone()
            if existing_by_date:
                bitacora_id = existing_by_date['id']
                cursor.execute('''
                    UPDATE bitacoras SET
                        user_id = COALESCE(?, user_id),
                        hora_inicio = ?,
                        colaborador = ?,
                        area = ?,
                        pendientes = ?,
                        necesita_apoyo = ?,
                        apoyo_detalle = ?,
                        prioridad_siguiente = ?,
                        tiempo_total_min = ?,
                        resumen_texto = ?,
                        estado = ?,
                        updated_at = ?
                    WHERE id = ?
                ''', (
                    user_id,
                    data.get('hora_inicio', ''),
                    colaborador,
                    area,
                    data.get('pendientes', ''),
                    data.get('necesita_apoyo', 'No'),
                    data.get('apoyo_detalle', ''),
                    data.get('prioridad_siguiente', ''),
                    tiempo_total,
                    data.get('resumen_texto', ''),
                    data.get('estado', 'generada'),
                    now_peru,
                    bitacora_id
                ))
            else:
                cursor.execute('''
                    INSERT INTO bitacoras (
                        user_id, fecha, hora_inicio, colaborador, area, pendientes,
                        necesita_apoyo, apoyo_detalle, prioridad_siguiente,
                        tiempo_total_min, resumen_texto, estado, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    user_id,
                    fecha,
                    data.get('hora_inicio', ''),
                    colaborador,
                    area,
                    data.get('pendientes', ''),
                    data.get('necesita_apoyo', 'No'),
                    data.get('apoyo_detalle', ''),
                    data.get('prioridad_siguiente', ''),
                    tiempo_total,
                    data.get('resumen_texto', ''),
                    data.get('estado', 'generada'),
                    now_peru,
                    now_peru
                ))
                bitacora_id = cursor.lastrowid


        # Mapeo de actividades previas para actualización quirúrgica (conserva IDs y evita duplicados al editar)
        existing_acts_list = [dict(r) for r in cursor.execute(
            "SELECT id, orden, hora_inicio, tipo_trabajo, descripcion, shared_uuid FROM actividades WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0) ORDER BY orden ASC, id ASC",
            (bitacora_id,)
        ).fetchall()]
        existing_acts_by_id = {r['id']: r for r in existing_acts_list}
        kept_ids = set()
        matched_target_ids = {}

        # Fase 1: Coincidencia exacta por ID numérico en la base de datos
        for idx, act in enumerate(actividades_data):
            act_id = act.get('id')
            act_id_num = None
            try:
                act_id_num = int(act_id)
            except (ValueError, TypeError):
                pass
            if act_id_num and act_id_num in existing_acts_by_id and act_id_num not in kept_ids:
                matched_target_ids[idx] = act_id_num
                kept_ids.add(act_id_num)

        # Fase 2: Coincidencia por shared_uuid para actividades sincronizadas entre usuarios
        for idx, act in enumerate(actividades_data):
            if idx in matched_target_ids:
                continue
            shared_uuid = act.get('shared_uuid')
            if shared_uuid:
                for ex in existing_acts_list:
                    if ex['id'] not in kept_ids and ex.get('shared_uuid') == shared_uuid:
                        matched_target_ids[idx] = ex['id']
                        kept_ids.add(ex['id'])
                        break

        # Fase 3: Coincidencia por posición (orden) para actividades que vienen sin ID numérico (ej. borrador local o ediciones en vuelo)
        # Esto PREVIENE que una edición de actividad genere un duplicado y mande la original a la papelera
        for idx, act in enumerate(actividades_data):
            if idx in matched_target_ids:
                continue
            act_orden = act.get('orden', idx)
            for ex in existing_acts_list:
                if ex['id'] not in kept_ids and ex.get('orden') == act_orden:
                    matched_target_ids[idx] = ex['id']
                    kept_ids.add(ex['id'])
                    break

        # Fase 4: Si la cantidad de actividades es >= a las existentes, emparejar unassigned por índice
        unmatched_existing = [ex['id'] for ex in existing_acts_list if ex['id'] not in kept_ids]
        unmatched_incoming = [idx for idx in range(len(actividades_data)) if idx not in matched_target_ids]
        if len(actividades_data) >= len(existing_acts_list):
            for i, idx in enumerate(unmatched_incoming):
                if i < len(unmatched_existing):
                    target_id = unmatched_existing[i]
                    matched_target_ids[idx] = target_id
                    kept_ids.add(target_id)

        for idx, act in enumerate(actividades_data):
            evidencias_val = act.get('evidencias', [])
            evidencias_json = json.dumps(evidencias_val) if isinstance(evidencias_val, (list, dict)) else (evidencias_val or '[]')
            
            shared_with_val = act.get('shared_with', [])
            if not isinstance(shared_with_val, list):
                try:
                    shared_with_val = json.loads(shared_with_val) if shared_with_val else []
                except Exception:
                    shared_with_val = []
            shared_with_json = json.dumps(shared_with_val)
            
            shared_uuid = act.get('shared_uuid')
            if shared_with_val and not shared_uuid:
                shared_uuid = f"sync-{int(time.time()*1000)}-{idx}-{user_id}"

            parent_task_id = int(act['parent_task_id']) if act.get('parent_task_id') and str(act['parent_task_id']).isdigit() else None
            comentarios = act.get('comentarios', '') or ''
            tipo_vinculo = act.get('tipo_vinculo', 'continuacion') or 'continuacion'

            target_act_id = matched_target_ids.get(idx)

            if target_act_id:
                cursor.execute('''
                    UPDATE actividades SET
                        orden = ?, hora_inicio = ?, duracion_min = ?,
                        tipo_trabajo = ?, descripcion = ?, para_cliente = ?,
                        estado = ?, evidencias = ?, shared_with = ?,
                        shared_uuid = COALESCE(?, shared_uuid),
                        comentarios = ?, parent_task_id = ?, tipo_vinculo = ?,
                        updated_at = ?
                    WHERE id = ?
                ''', (
                    idx,
                    act.get('hora_inicio', ''),
                    int(act.get('duracion_min', 0) or 0),
                    act.get('tipo_trabajo', ''),
                    act.get('descripcion', ''),
                    act.get('para_cliente', ''),
                    act.get('estado', 'completada'),
                    evidencias_json,
                    shared_with_json,
                    shared_uuid,
                    comentarios,
                    parent_task_id,
                    tipo_vinculo,
                    now_peru,
                    target_act_id
                ))
                kept_ids.add(target_act_id)
            else:
                cursor.execute('''
                    INSERT INTO actividades (
                        bitacora_id, orden, hora_inicio, duracion_min,
                        tipo_trabajo, descripcion, para_cliente, estado, evidencias,
                        shared_with, shared_uuid, comentarios, parent_task_id, tipo_vinculo,
                        created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    bitacora_id,
                    idx,
                    act.get('hora_inicio', ''),
                    int(act.get('duracion_min', 0) or 0),
                    act.get('tipo_trabajo', ''),
                    act.get('descripcion', ''),
                    act.get('para_cliente', ''),
                    act.get('estado', 'completada'),
                    evidencias_json,
                    shared_with_json,
                    shared_uuid,
                    comentarios,
                    parent_task_id,
                    tipo_vinculo,
                    now_peru,
                    now_peru
                ))
                kept_ids.add(cursor.lastrowid)

            # Sincronización automática de tarea compartida en paralelo con el otro usuario
            if shared_with_val and shared_uuid:
                for target_uid in shared_with_val:
                    if target_uid != user_id:
                        target_u = cursor.execute("SELECT id, full_name FROM users WHERE id = ?", (target_uid,)).fetchone()
                        if target_u:
                            target_b = cursor.execute(
                                "SELECT id FROM bitacoras WHERE (user_id = ? OR LOWER(colaborador) = LOWER(?)) AND fecha = ?",
                                (target_u['id'], target_u['full_name'], fecha)
                            ).fetchone()
                            if not target_b:
                                cursor.execute('''
                                    INSERT INTO bitacoras (
                                        user_id, fecha, hora_inicio, colaborador, area,
                                        pendientes, necesita_apoyo, prioridad_siguiente,
                                        tiempo_total_min, estado, created_at, updated_at
                                    ) VALUES (?, ?, ?, ?, 'Sistemas', '', 'No', '', 0, 'generada', ?, ?)
                                ''', (target_u['id'], fecha, act.get('hora_inicio') or get_default_hora_inicio(), target_u['full_name'], now_peru, now_peru))
                                target_b_id = cursor.lastrowid
                            else:
                                target_b_id = target_b['id']

                            existing_sync_act = cursor.execute(
                                "SELECT id FROM actividades WHERE bitacora_id = ? AND shared_uuid = ?",
                                (target_b_id, shared_uuid)
                            ).fetchone()
                            
                            counterpart_shared = json.dumps([user_id])
                            
                            if existing_sync_act:
                                cursor.execute('''
                                    UPDATE actividades SET
                                        hora_inicio = ?,
                                        duracion_min = ?,
                                        tipo_trabajo = ?,
                                        descripcion = ?,
                                        para_cliente = ?,
                                        estado = ?,
                                        evidencias = ?,
                                        shared_with = ?,
                                        updated_at = ?
                                    WHERE id = ?
                                ''', (
                                    act.get('hora_inicio', ''),
                                    int(act.get('duracion_min', 0) or 0),
                                    act.get('tipo_trabajo', ''),
                                    act.get('descripcion', ''),
                                    act.get('para_cliente', ''),
                                    act.get('estado', 'en_proceso'),
                                    evidencias_json,
                                    counterpart_shared,
                                    now_peru,
                                    existing_sync_act['id']
                                ))
                            else:
                                next_ord = cursor.execute(
                                    "SELECT COALESCE(MAX(orden), 0) + 1 FROM actividades WHERE bitacora_id = ?",
                                    (target_b_id,)
                                ).fetchone()[0]
                                cursor.execute('''
                                    INSERT INTO actividades (
                                        bitacora_id, orden, hora_inicio, duracion_min,
                                        tipo_trabajo, descripcion, para_cliente, estado,
                                        evidencias, shared_with, shared_uuid, created_at, updated_at
                                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                ''', (
                                    target_b_id,
                                    next_ord,
                                    act.get('hora_inicio', ''),
                                    int(act.get('duracion_min', 0) or 0),
                                    act.get('tipo_trabajo', ''),
                                    act.get('descripcion', ''),
                                    act.get('para_cliente', ''),
                                    act.get('estado', 'en_proceso'),
                                    evidencias_json,
                                    counterpart_shared,
                                    shared_uuid,
                                    now_peru,
                                    now_peru
                                ))
                                
                            # Recalcular el tiempo_total_min de la bitácora del colaborador
                            cursor.execute('''
                                UPDATE bitacoras SET tiempo_total_min = (
                                    SELECT COALESCE(SUM(duracion_min), 0) FROM actividades WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)
                                ), updated_at = ? WHERE id = ?
                            ''', (now_peru, target_b_id, target_b_id))

        # Safeguard de integridad: si vienen 0 actividades pero la bitácora ya tenía actividades previas,
        # NO borrarlas todas a menos que venga explícitamente confirm_empty=True
        # Esto previene pérdida accidental de datos por fallas de red, debounce o cambios rápidos de fecha
        if len(actividades_data) == 0 and len(existing_acts_by_id) > 0 and not data.get('confirm_empty'):
            kept_ids = set(existing_acts_by_id.keys())

        # Mover a papelera únicamente actividades removidas conscientemente por el usuario
        for old_id in existing_acts_by_id:
            if old_id not in kept_ids:
                cursor.execute("UPDATE actividades SET parent_task_id = NULL WHERE parent_task_id = ?", (old_id,))
                cursor.execute("UPDATE actividades SET is_deleted = 1, deleted_at = ? WHERE id = ?", (now_peru, old_id))


        conn.commit()
        
        users_map = {u['id']: u['full_name'] for u in cursor.execute("SELECT id, full_name FROM users").fetchall()}
        b = cursor.execute("SELECT * FROM bitacoras WHERE id = ?", (bitacora_id,)).fetchone()
        b_dict = dict(b)
        acts = cursor.execute('''
            SELECT a.*, p.descripcion as parent_task_desc, p.estado as parent_task_estado 
            FROM actividades a 
            LEFT JOIN actividades p ON a.parent_task_id = p.id 
            WHERE a.bitacora_id = ? AND (a.is_deleted IS NULL OR a.is_deleted = 0) 
            ORDER BY a.orden ASC, a.id ASC
        ''', (bitacora_id,)).fetchall()
        b_dict['actividades'] = [format_actividad_dict(a, users_map) for a in acts]
        
        conn.close()
        return jsonify({"message": "Bitácora guardada con éxito", "bitacora": b_dict}), 200
        
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({"error": str(e)}), 500

@app.route('/api/actividades/<int:act_id>/estado', methods=['PATCH', 'OPTIONS'])
def update_actividad_estado(act_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    data = request.get_json() or {}
    nuevo_estado = data.get('estado')
    if not nuevo_estado:
        return jsonify({"error": "Estado es requerido"}), 400
        
    conn = get_db()
    cursor = conn.cursor()
    now_peru = get_peru_now_str()
    act = cursor.execute("SELECT shared_uuid FROM actividades WHERE id = ?", (act_id,)).fetchone()
    if act and act['shared_uuid']:
        cursor.execute("UPDATE actividades SET estado = ?, updated_at = ? WHERE shared_uuid = ?", (nuevo_estado, now_peru, act['shared_uuid']))
    else:
        cursor.execute("UPDATE actividades SET estado = ?, updated_at = ? WHERE id = ?", (nuevo_estado, now_peru, act_id))
    conn.commit()
    conn.close()
    return jsonify({"success": True, "id": act_id, "nuevo_estado": nuevo_estado})

@app.route('/api/actividades/<int:act_id>', methods=['GET', 'PUT', 'OPTIONS'])
def update_actividad_detalle(act_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    if request.method == 'GET':
        conn = get_db()
        cursor = conn.cursor()
        query = '''
            SELECT a.*, b.fecha as bitacora_fecha, b.colaborador
            FROM actividades a
            JOIN bitacoras b ON a.bitacora_id = b.id
            WHERE a.id = ? AND (a.is_deleted IS NULL OR a.is_deleted = 0)
        '''
        row = cursor.execute(query, (act_id,)).fetchone()
        conn.close()
        if not row:
            return jsonify({'error': 'Actividad no encontrada'}), 404
        d = dict(row)
        if d.get('evidencias'):
            try:
                d['evidencias'] = json.loads(d['evidencias'])
            except Exception:
                pass
        if d.get('shared_with'):
            try:
                d['shared_with'] = json.loads(d['shared_with'])
            except Exception:
                pass
        return jsonify(d), 200

    data = request.get_json() or {}
    
    evidencias_json = None
    if 'evidencias' in data:
        ev_val = data['evidencias']
        evidencias_json = json.dumps(ev_val) if isinstance(ev_val, (list, dict)) else str(ev_val)
        
    conn = get_db()
    cursor = conn.cursor()
    now_peru = get_peru_now_str()
    act = cursor.execute("SELECT shared_uuid FROM actividades WHERE id = ?", (act_id,)).fetchone()
    
    shared_with_val = data.get('shared_with')
    shared_with_json = json.dumps(shared_with_val) if isinstance(shared_with_val, list) else None

    parent_task_id_param = int(data['parent_task_id']) if data.get('parent_task_id') and str(data['parent_task_id']).isdigit() else (None if 'parent_task_id' in data else None)
    comentarios_param = data.get('comentarios')
    tipo_vinculo_param = data.get('tipo_vinculo')

    if act and act['shared_uuid']:
        cursor.execute('''
            UPDATE actividades SET
                hora_inicio = COALESCE(?, hora_inicio),
                duracion_min = COALESCE(?, duracion_min),
                tipo_trabajo = COALESCE(?, tipo_trabajo),
                descripcion = COALESCE(?, descripcion),
                para_cliente = COALESCE(?, para_cliente),
                estado = COALESCE(?, estado),
                evidencias = COALESCE(?, evidencias),
                shared_with = COALESCE(?, shared_with),
                comentarios = COALESCE(?, comentarios),
                parent_task_id = COALESCE(?, parent_task_id),
                tipo_vinculo = COALESCE(?, tipo_vinculo),
                updated_at = ?
            WHERE shared_uuid = ?
        ''', (
            data.get('hora_inicio'),
            data.get('duracion_min'),
            data.get('tipo_trabajo'),
            data.get('descripcion'),
            data.get('para_cliente'),
            data.get('estado'),
            evidencias_json,
            shared_with_json,
            comentarios_param,
            parent_task_id_param,
            tipo_vinculo_param,
            now_peru,
            act['shared_uuid']
        ))
    else:
        cursor.execute('''
            UPDATE actividades SET
                hora_inicio = COALESCE(?, hora_inicio),
                duracion_min = COALESCE(?, duracion_min),
                tipo_trabajo = COALESCE(?, tipo_trabajo),
                descripcion = COALESCE(?, descripcion),
                para_cliente = COALESCE(?, para_cliente),
                estado = COALESCE(?, estado),
                evidencias = COALESCE(?, evidencias),
                shared_with = COALESCE(?, shared_with),
                comentarios = COALESCE(?, comentarios),
                parent_task_id = COALESCE(?, parent_task_id),
                tipo_vinculo = COALESCE(?, tipo_vinculo),
                updated_at = ?
            WHERE id = ?
        ''', (
            data.get('hora_inicio'),
            data.get('duracion_min'),
            data.get('tipo_trabajo'),
            data.get('descripcion'),
            data.get('para_cliente'),
            data.get('estado'),
            evidencias_json,
            shared_with_json,
            comentarios_param,
            parent_task_id_param,
            tipo_vinculo_param,
            now_peru,
            act_id
        ))
        
    # Recalcular el tiempo_total_min de la bitácora asociada
    b_row = cursor.execute("SELECT bitacora_id FROM actividades WHERE id = ?", (act_id,)).fetchone()
    if b_row and b_row['bitacora_id']:
        cursor.execute('''
            UPDATE bitacoras SET tiempo_total_min = (
                SELECT COALESCE(SUM(duracion_min), 0) FROM actividades 
                WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)
            ), updated_at = ? WHERE id = ?
        ''', (b_row['bitacora_id'], now_peru, b_row['bitacora_id']))

    conn.commit()
    users_map = {u['id']: u['full_name'] for u in cursor.execute("SELECT id, full_name FROM users").fetchall()}
    updated_act = cursor.execute('''
        SELECT a.*, p.descripcion as parent_task_desc, p.estado as parent_task_estado
        FROM actividades a
        LEFT JOIN actividades p ON a.parent_task_id = p.id
        WHERE a.id = ?
    ''', (act_id,)).fetchone()
    conn.close()
    return jsonify({"success": True, "actividad": format_actividad_dict(updated_act, users_map) if updated_act else None})


@app.route('/api/actividades/pendientes', methods=['GET'])
def get_actividades_pendientes():
    colaborador = request.args.get('colaborador')
    user_id = request.args.get('user_id')
    
    conn = get_db()
    cursor = conn.cursor()
    
    query = '''
        SELECT a.*, b.fecha, b.colaborador
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        WHERE a.estado != 'completada' AND (a.is_deleted IS NULL OR a.is_deleted = 0)
    '''
    params = []
    if colaborador:
        query += " AND b.colaborador LIKE ?"
        params.append(f"%{colaborador}%")
    if user_id:
        query += " AND b.user_id = ?"
        params.append(user_id)
        
    query += " ORDER BY a.id DESC LIMIT 50"
    acts = cursor.execute(query, params).fetchall()
    conn.close()
    return jsonify([format_actividad_dict(a) for a in acts])

@app.route('/api/bitacoras/<int:bitacora_id>', methods=['DELETE', 'OPTIONS'])
def delete_bitacora(bitacora_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE actividades SET parent_task_id = NULL WHERE parent_task_id IN (SELECT id FROM actividades WHERE bitacora_id = ?)", (bitacora_id,))
    cursor.execute("DELETE FROM actividades WHERE bitacora_id = ?", (bitacora_id,))
    cursor.execute("DELETE FROM bitacoras WHERE id = ?", (bitacora_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True, "message": "Bitácora eliminada"})

# ==================== PAPELERA DE RECICLAJE (RETENCIÓN 15 DÍAS) ====================

@app.route('/api/papelera', methods=['GET', 'OPTIONS'])
def get_papelera():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    
    conn = get_db()
    try:
        cursor = conn.cursor()
        user_id = request.args.get('user_id')
        requesting_user_id = request.args.get('requesting_user_id')
        
        query = '''
            SELECT a.*, b.fecha as bitacora_fecha, b.colaborador as bitacora_colaborador, b.user_id as bitacora_user_id,
                   u.team_id, t.nombre as team_name
            FROM actividades a
            JOIN bitacoras b ON a.bitacora_id = b.id
            LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE LOWER(full_name) = LOWER(b.colaborador) OR LOWER(username) = LOWER(b.colaborador) LIMIT 1)) = u.id
            LEFT JOIN teams t ON u.team_id = t.id
            WHERE a.is_deleted = 1
        '''
        params = []
        
        if requesting_user_id:
            req_u = cursor.execute("SELECT id, role, team_id, full_name, username FROM users WHERE id = ?", (requesting_user_id,)).fetchone()
            if req_u:
                if req_u['role'] in ('analista', 'operador'):
                    query += " AND (b.user_id = ? OR LOWER(b.colaborador) = LOWER(?) OR LOWER(b.colaborador) = LOWER(?))"
                    params.extend([req_u['id'], req_u['full_name'], req_u['username']])
                # Admin y líderes pueden supervisar la papelera del equipo
        elif user_id:
            query += " AND (b.user_id = ? OR u.id = ?)"
            params.extend([user_id, user_id])
            
        query += " ORDER BY a.deleted_at DESC, a.id DESC"
        deleted_rows = cursor.execute(query, params).fetchall()
        users_map = {u['id']: u['full_name'] for u in cursor.execute("SELECT id, full_name FROM users").fetchall()}
        
        items = []

        for row in deleted_rows:
            item = format_actividad_dict(row, users_map)
            deleted_at_str = item.get('deleted_at')
            dias_restantes = 15
            expira_en = None
            
            if deleted_at_str:
                try:
                    del_dt = datetime.strptime(deleted_at_str.split('.')[0], '%Y-%m-%d %H:%M:%S')
                    now_p = get_peru_now()
                    now_naive = datetime(now_p.year, now_p.month, now_p.day, now_p.hour, now_p.minute, now_p.second)
                    dias_transcurridos = max(0, (now_naive - del_dt).days)
                    dias_restantes = max(0, min(15, 15 - dias_transcurridos))
                    expira_dt = del_dt + timedelta(days=15)
                    expira_en = expira_dt.strftime('%Y-%m-%d')
                except Exception:
                    dias_restantes = 15
                    
            item['dias_restantes'] = dias_restantes
            item['expira_en'] = expira_en
            items.append(item)
            
        return jsonify({
            "success": True,
            "items": items,
            "total": len(items)
        })
    finally:
        conn.close()

@app.route('/api/papelera/restaurar/<int:act_id>', methods=['POST', 'OPTIONS'])
def restaurar_actividad(act_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    data = request.get_json(silent=True) or {}
    target_bitacora_id = data.get('target_bitacora_id')
    target_fecha = data.get('target_fecha')
    target_user_id = data.get('target_user_id')
        
    conn = get_db()
    cursor = conn.cursor()
    now_peru = get_peru_now_str()
    
    act = cursor.execute("SELECT * FROM actividades WHERE id = ?", (act_id,)).fetchone()
    if not act:
        conn.close()
        return jsonify({"error": "Actividad no encontrada"}), 404
        
    old_bitacora_id = act['bitacora_id']
    new_bitacora_id = old_bitacora_id

    # Si se especificó una bitácora destino puntual
    if target_bitacora_id and target_bitacora_id != old_bitacora_id:
        target_b = cursor.execute("SELECT id FROM bitacoras WHERE id = ?", (target_bitacora_id,)).fetchone()
        if target_b:
            new_bitacora_id = target_bitacora_id
    # O si se especificó una fecha destino (ej. hoy) para un usuario
    elif target_fecha and target_user_id:
        existing_b = cursor.execute("SELECT id FROM bitacoras WHERE user_id = ? AND fecha = ?", (target_user_id, target_fecha)).fetchone()
        if existing_b:
            new_bitacora_id = existing_b['id']

    # Restaurar actividad
    cursor.execute('''
        UPDATE actividades 
        SET is_deleted = 0, deleted_at = NULL, bitacora_id = ?, updated_at = ? 
        WHERE id = ?
    ''', (new_bitacora_id, now_peru, act_id))
    
    # Recalcular tiempo total de la bitácora antigua
    cursor.execute('''
        UPDATE bitacoras SET tiempo_total_min = (
            SELECT COALESCE(SUM(duracion_min), 0) FROM actividades 
            WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)
        ), updated_at = ? WHERE id = ?
    ''', (old_bitacora_id, now_peru, old_bitacora_id))

    # Si cambió de bitácora, recalcular también la nueva
    if new_bitacora_id != old_bitacora_id:
        cursor.execute('''
            UPDATE bitacoras SET tiempo_total_min = (
                SELECT COALESCE(SUM(duracion_min), 0) FROM actividades 
                WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)
            ), updated_at = ? WHERE id = ?
        ''', (new_bitacora_id, now_peru, new_bitacora_id))
    
    conn.commit()

    restored_row = cursor.execute('''
        SELECT a.*, b.fecha as bitacora_fecha, b.colaborador as bitacora_colaborador
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        WHERE a.id = ?
    ''', (act_id,)).fetchone()
    
    users_map = {u['id']: u['full_name'] for u in cursor.execute("SELECT id, full_name FROM users").fetchall()}
    formatted_act = format_actividad_dict(restored_row, users_map) if restored_row else dict(act)

    conn.close()
    
    return jsonify({
        "success": True, 
        "message": "Actividad restaurada exitosamente", 
        "id": act_id,
        "actividad": formatted_act,
        "bitacora_id": new_bitacora_id,
        "bitacora_fecha": restored_row['bitacora_fecha'] if restored_row else None
    })

@app.route('/api/papelera/eliminar/<int:act_id>', methods=['DELETE', 'OPTIONS'])
def eliminar_definitivo_actividad(act_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE actividades SET parent_task_id = NULL WHERE parent_task_id = ?", (act_id,))
    cursor.execute("DELETE FROM actividades WHERE id = ?", (act_id,))
    conn.commit()
    conn.close()
    
    return jsonify({"success": True, "message": "Actividad eliminada definitivamente"})

@app.route('/api/papelera/vaciar', methods=['POST', 'OPTIONS'])
def vaciar_papelera():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    data = request.get_json() or {}
    user_id = data.get('user_id')
    role = data.get('role')
    
    conn = get_db()
    cursor = conn.cursor()
    
    if role == 'admin' and not user_id:
        cursor.execute("UPDATE actividades SET parent_task_id = NULL WHERE parent_task_id IN (SELECT id FROM actividades WHERE is_deleted = 1)")
        cursor.execute("DELETE FROM actividades WHERE is_deleted = 1")
    elif user_id:
        cursor.execute('''
            UPDATE actividades SET parent_task_id = NULL 
            WHERE parent_task_id IN (
                SELECT id FROM actividades 
                WHERE is_deleted = 1 AND bitacora_id IN (
                    SELECT id FROM bitacoras WHERE user_id = ? OR colaborador = (SELECT full_name FROM users WHERE id = ?)
                )
            )
        ''', (user_id, user_id))
        cursor.execute('''
            DELETE FROM actividades 
            WHERE is_deleted = 1 AND bitacora_id IN (
                SELECT id FROM bitacoras WHERE user_id = ? OR colaborador = (SELECT full_name FROM users WHERE id = ?)
            )
        ''', (user_id, user_id))
    else:
        cursor.execute("UPDATE actividades SET parent_task_id = NULL WHERE parent_task_id IN (SELECT id FROM actividades WHERE is_deleted = 1)")
        cursor.execute("DELETE FROM actividades WHERE is_deleted = 1")
        
    conn.commit()
    conn.close()
    return jsonify({"success": True, "message": "Papelera vaciada correctamente"})

@app.route('/api/actividades/<int:act_id>/eliminar', methods=['POST', 'DELETE', 'OPTIONS'])
def soft_delete_actividad(act_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    act = cursor.execute("SELECT bitacora_id FROM actividades WHERE id = ?", (act_id,)).fetchone()
    if not act:
        conn.close()
        return jsonify({"error": "Actividad no encontrada"}), 404
        
    now_peru = get_peru_now_str()
    cursor.execute("UPDATE actividades SET is_deleted = 1, deleted_at = ? WHERE id = ?", (now_peru, act_id))
    
    # Recalcular tiempo_total_min
    cursor.execute('''
        UPDATE bitacoras SET tiempo_total_min = (
            SELECT COALESCE(SUM(duracion_min), 0) FROM actividades 
            WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)
        ), updated_at = ? WHERE id = ?
    ''', (act['bitacora_id'], now_peru, act['bitacora_id']))

    
    conn.commit()
    conn.close()
    return jsonify({"success": True, "message": "Actividad movida a la papelera (retención 15 días)"})

# ==================== SUPERVISIÓN EN VIVO: FEED DE EQUIPO ====================

@app.route('/api/equipo/actividades-en-vivo', methods=['GET', 'OPTIONS'])
def get_equipo_actividades_en_vivo():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
        
    conn = get_db()
    cursor = conn.cursor()
    
    fecha = request.args.get('fecha') or get_peru_today_str()
    requesting_user_id = request.args.get('requesting_user_id')
    team_id = request.args.get('team_id')
    
    query = '''
        SELECT a.*, b.fecha as bitacora_fecha, b.colaborador, b.user_id, b.area,
               b.necesita_apoyo, b.apoyo_detalle,
               u.team_id, t.nombre as team_name, u.full_name as user_full_name, u.phone as colaborador_phone
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE LOWER(full_name) = LOWER(b.colaborador) OR LOWER(username) = LOWER(b.colaborador) LIMIT 1)) = u.id
        LEFT JOIN teams t ON u.team_id = t.id
        WHERE b.fecha = ? AND (a.is_deleted IS NULL OR a.is_deleted = 0)
    '''
    params = [fecha]
    
    if requesting_user_id:
        req_u = cursor.execute("SELECT id, role, team_id, full_name, username FROM users WHERE id = ?", (requesting_user_id,)).fetchone()
        if req_u:
            if req_u['role'] in ('analista', 'operador'):
                query += " AND (b.user_id = ? OR LOWER(b.colaborador) = LOWER(?) OR LOWER(b.colaborador) = LOWER(?))"
                params.extend([req_u['id'], req_u['full_name'], req_u['username']])
            # Admin y líderes pueden ver las actividades de todo el equipo o filtrar por team_id
                    
    if team_id and team_id != 'all':
        query += " AND u.team_id = ?"
        params.append(team_id)
        
    query += " ORDER BY a.updated_at DESC, a.id DESC"
    rows = cursor.execute(query, params).fetchall()
    users_map = {u['id']: u['full_name'] for u in cursor.execute("SELECT id, full_name FROM users").fetchall()}
    
    feed = [format_actividad_dict(r, users_map) for r in rows]
    conn.close()
    
    return jsonify({
        "success": True,
        "fecha": fecha,
        "total": len(feed),
        "actividades": feed
    })

# ==================== DASHBOARD & ANALYTICS ====================

@app.route('/api/dashboard/stats', methods=['GET'])
def get_dashboard_stats():
    team_id = request.args.get('team_id')
    user_id = request.args.get('user_id')
    fecha = request.args.get('fecha')
    requesting_user_id = request.args.get('requesting_user_id')
    
    conn = get_db()
    cursor = conn.cursor()

    # Validación de privacidad por rol
    if requesting_user_id:
        req_u = cursor.execute("SELECT id, role, team_id FROM users WHERE id = ?", (requesting_user_id,)).fetchone()
        if req_u:
            if req_u['role'] in ('analista', 'operador'):
                user_id = req_u['id']
                team_id = None
            elif req_u['role'] == 'lider':
                if req_u['team_id'] and not team_id:
                    team_id = req_u['team_id']
    
    base_filter = "WHERE 1=1"
    params = []
    
    if team_id:
        base_filter += " AND u.team_id = ?"
        params.append(team_id)
    if user_id:
        base_filter += " AND (b.user_id = ? OR u.id = ?)"
        params.extend([user_id, user_id])
    if fecha:
        base_filter += " AND b.fecha = ?"
        params.append(fecha)
        
    # 1. Totales generales
    total_query = f'''
        SELECT COUNT(DISTINCT b.id) as total_bitacoras,
               COALESCE(SUM(b.tiempo_total_min), 0) as total_minutos,
               COUNT(DISTINCT b.colaborador) as total_colaboradores
        FROM bitacoras b
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        {base_filter}
    '''
    totales = cursor.execute(total_query, params).fetchone()
    
    # 2. Desglose de actividades por estado
    act_filter_query = f'''
        SELECT a.estado, COUNT(a.id) as cantidad, COALESCE(SUM(a.duracion_min), 0) as minutos
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        {base_filter} AND (a.is_deleted IS NULL OR a.is_deleted = 0)
        GROUP BY a.estado
    '''
    act_stats = cursor.execute(act_filter_query, params).fetchall()
    
    # 3. Desglose por tipo de trabajo
    tipo_query = f'''
        SELECT COALESCE(NULLIF(a.tipo_trabajo, ''), 'No clasificado') as tipo,
               COUNT(a.id) as cantidad,
               COALESCE(SUM(a.duracion_min), 0) as minutos
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        {base_filter} AND (a.is_deleted IS NULL OR a.is_deleted = 0)
        GROUP BY tipo
        ORDER BY minutos DESC
    '''
    tipo_stats = cursor.execute(tipo_query, params).fetchall()

    # 3.1 Desglose por cliente / proyecto
    cliente_query = f'''
        SELECT COALESCE(NULLIF(a.para_cliente, ''), 'General / Interno') as cliente,
               COUNT(a.id) as cantidad,
               COALESCE(SUM(a.duracion_min), 0) as minutos
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        {base_filter} AND (a.is_deleted IS NULL OR a.is_deleted = 0)
        GROUP BY cliente
        ORDER BY minutos DESC
        LIMIT 10
    '''
    cliente_stats = cursor.execute(cliente_query, params).fetchall()
    
    # 4. Desglose por colaborador (calcula duración real desde actividades y evita inflación cartesiana)
    colab_query = f'''
        SELECT b.colaborador,
               u.id as user_id,
               u.team_id,
               t.nombre as team_name,
               COUNT(DISTINCT b.id) as bitacoras_count,
               COALESCE(SUM(a.duracion_min), 0) as total_minutos,
               COUNT(a.id) as total_actividades,
               SUM(CASE WHEN a.estado = 'completada' THEN 1 ELSE 0 END) as actividades_completadas
        FROM bitacoras b
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        LEFT JOIN teams t ON u.team_id = t.id
        LEFT JOIN actividades a ON a.bitacora_id = b.id AND (a.is_deleted IS NULL OR a.is_deleted = 0)
        {base_filter}
        GROUP BY b.colaborador
        ORDER BY total_minutos DESC
    '''
    colab_stats = cursor.execute(colab_query, params).fetchall()
    
    # 5. Colaboradores que requieren apoyo / bloqueos activos
    apoyo_query = f'''
        SELECT b.id, b.fecha, b.colaborador, b.apoyo_detalle, b.pendientes, t.nombre as team_name
        FROM bitacoras b
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        LEFT JOIN teams t ON u.team_id = t.id
        {base_filter} AND b.necesita_apoyo = 'Si'
        ORDER BY b.fecha DESC, b.id DESC
        LIMIT 10
    '''
    alertas_apoyo = cursor.execute(apoyo_query, params).fetchall()

    # 6. Tendencia cronológica diaria (últimas fechas para gráfico Spline/Líneas)
    tendencia_query = f'''
        SELECT b.fecha,
               COUNT(a.id) as actividades,
               COALESCE(SUM(a.duracion_min), 0) as minutos,
               SUM(CASE WHEN a.estado = 'completada' THEN 1 ELSE 0 END) as completadas
        FROM bitacoras b
        LEFT JOIN actividades a ON a.bitacora_id = b.id AND (a.is_deleted IS NULL OR a.is_deleted = 0)
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        {base_filter}
        GROUP BY b.fecha
        ORDER BY b.fecha ASC
        LIMIT 30
    '''
    tendencia_diaria = cursor.execute(tendencia_query, params).fetchall()

    # 7. Distribución por franja horaria del día
    horaria_query = f'''
        SELECT 
            CASE 
                WHEN a.hora_inicio < '10:00' THEN '08:00 - 10:00'
                WHEN a.hora_inicio < '12:00' THEN '10:00 - 12:00'
                WHEN a.hora_inicio < '14:00' THEN '12:00 - 14:00'
                WHEN a.hora_inicio < '16:00' THEN '14:00 - 16:00'
                WHEN a.hora_inicio < '18:00' THEN '16:00 - 18:00'
                ELSE '18:00+'
            END as franja,
            COUNT(a.id) as cantidad,
            COALESCE(SUM(a.duracion_min), 0) as minutos
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
        {base_filter} AND (a.is_deleted IS NULL OR a.is_deleted = 0) AND a.hora_inicio IS NOT NULL AND a.hora_inicio != ''
        GROUP BY franja
        ORDER BY MIN(a.hora_inicio) ASC
    '''
    distribucion_horaria = cursor.execute(horaria_query, params).fetchall()

    # 8. Métricas y KPIs de telemetría operativa
    tot_min = totales['total_minutos'] if totales else 0
    tot_bit = totales['total_bitacoras'] if totales else 0
    promedio_jornada = round(tot_min / tot_bit) if tot_bit > 0 else 0

    total_acts = sum(r['cantidad'] for r in act_stats) if act_stats else 0
    completed_acts = sum(r['cantidad'] for r in act_stats if r['estado'] == 'completada') if act_stats else 0
    eficiencia = round((completed_acts / total_acts) * 100) if total_acts > 0 else 0

    planificadas_min = sum(r['minutos'] for r in tipo_stats if 'planificad' in (r['tipo'] or '').lower()) if tipo_stats else 0
    ratio_planificado = round((planificadas_min / tot_min) * 100) if tot_min > 0 else 0

    # 9. Historial reciente de bitácoras del usuario (para tabla individual de operador)
    jornadas_usuario = []
    if user_id:
        jornadas_query = f'''
            SELECT b.id, b.fecha, b.tiempo_total_min, b.necesita_apoyo, b.apoyo_detalle,
                   COUNT(a.id) as total_actividades,
                   SUM(CASE WHEN a.estado = 'completada' THEN 1 ELSE 0 END) as actividades_completadas
            FROM bitacoras b
            LEFT JOIN actividades a ON a.bitacora_id = b.id AND (a.is_deleted IS NULL OR a.is_deleted = 0)
            LEFT JOIN users u ON COALESCE(b.user_id, (SELECT id FROM users WHERE full_name = b.colaborador LIMIT 1)) = u.id
            {base_filter}
            GROUP BY b.id
            ORDER BY b.fecha DESC
            LIMIT 15
        '''
        jornadas_usuario = cursor.execute(jornadas_query, params).fetchall()
    
    conn.close()
    
    return jsonify({
        "resumen": dict(totales) if totales else {},
        "por_estado": [dict(r) for r in act_stats],
        "por_tipo": [dict(r) for r in tipo_stats],
        "por_cliente": [dict(r) for r in cliente_stats],
        "por_colaborador": [dict(r) for r in colab_stats],
        "alertas_apoyo": [dict(r) for r in alertas_apoyo],
        "tendencia_diaria": [dict(r) for r in tendencia_diaria],
        "distribucion_horaria": [dict(r) for r in distribucion_horaria],
        "jornadas_recientes": [dict(r) for r in jornadas_usuario],
        "kpis_adicionales": {
            "promedio_minutos_jornada": promedio_jornada,
            "eficiencia": eficiencia,
            "ratio_planificado": ratio_planificado,
            "total_actividades": total_acts,
            "actividades_completadas": completed_acts
        }
    })

# ==================== CRON ROLLOVER & TASK LIFECYCLE ====================

def execute_daily_rollover(today_str=None):
    """
    Motor automático de medianoche (00:00) y catch-up:
    1. Cierra automáticamente bitácoras de días anteriores que quedaron abiertas ('cerrada_sistema').
    2. Identifica actividades no completadas ('pendiente', 'en_proceso', 'en_revision') o modificadas
       y las transfiere a la jornada de hoy con duracion_min = 0 para no duplicar horas.
    3. Auto-purga actividades en papelera que tengan más de 15 días.
    """
    if not today_str:
        today_str = get_peru_today_str()
        
    conn = get_db()
    cursor = conn.cursor()
    closed_count = 0
    rolled_count = 0
    now_peru = get_peru_now_str()
    
    try:
        # 0. Auto-purga permanente de actividades eliminadas con más de 15 días de antigüedad según horario de Perú
        purge_threshold = (get_peru_now() - timedelta(days=15)).strftime('%Y-%m-%d %H:%M:%S')
        cursor.execute("UPDATE actividades SET parent_task_id = NULL WHERE parent_task_id IN (SELECT id FROM actividades WHERE is_deleted = 1 AND deleted_at < ?)", (purge_threshold,))
        cursor.execute("DELETE FROM actividades WHERE is_deleted = 1 AND deleted_at < ?", (purge_threshold,))

        # 1. Auto-cerrar bitácoras de días pasados (< today_str) que no estén formalmente cerradas
        past_open = cursor.execute('''
            SELECT id, user_id, colaborador, fecha, pendientes
            FROM bitacoras
            WHERE fecha < ? AND estado NOT IN ('cerrada', 'cerrada_sistema')
        ''', (today_str,)).fetchall()
        
        for b in past_open:
            dur = cursor.execute(
                "SELECT COALESCE(SUM(duracion_min), 0) FROM actividades WHERE bitacora_id = ? AND (is_deleted IS NULL OR is_deleted = 0)", 
                (b['id'],)
            ).fetchone()[0]
            incomp = cursor.execute(
                "SELECT descripcion FROM actividades WHERE bitacora_id = ? AND estado != 'completada' AND (is_deleted IS NULL OR is_deleted = 0)", 
                (b['id'],)
            ).fetchall()
            auto_pend = b['pendientes']
            if not auto_pend or len(auto_pend.strip()) == 0:
                if incomp:
                    auto_pend = "Pendientes al cierre: " + "; ".join(r['descripcion'] for r in incomp[:3])
                else:
                    auto_pend = "Todas las actividades fueron concluidas en esta fecha."
                    
            cursor.execute('''
                UPDATE bitacoras SET
                    estado = 'cerrada_sistema',
                    tiempo_total_min = ?,
                    pendientes = ?,
                    updated_at = ?
                WHERE id = ?
            ''', (dur, auto_pend, now_peru, b['id']))
            closed_count += 1

        # 2. Arrastre de actividades abiertas hacia el día de hoy
        active_users = cursor.execute("SELECT id, full_name FROM users WHERE is_active = 1").fetchall()
        for u in active_users:
            rolled_count += rollover_user_open_tasks(cursor, u['id'], u['full_name'], today_str, now_peru)
                    
        conn.commit()
    except Exception as e:
        conn.rollback()
        print("Error en execute_daily_rollover:", e)
    finally:
        conn.close()
        
    return {
        "success": True,
        "date": today_str,
        "closed_bitacoras": closed_count,
        "rolled_activities": rolled_count
    }

@app.route('/api/cron/rollover', methods=['GET', 'POST', 'OPTIONS'])
def trigger_rollover():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    date_arg = request.args.get('date')
    res = execute_daily_rollover(today_str=date_arg)
    return jsonify(res)

@app.route('/api/actividades/importar-pendientes', methods=['POST', 'OPTIONS'])
def importar_actividades_pendientes():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    data = request.get_json() or {}
    colaborador = data.get('colaborador')
    user_id = data.get('user_id')
    target_fecha = data.get('target_fecha') or get_peru_today_str()
    
    if not colaborador:
        return jsonify({"error": "Colaborador es requerido"}), 400
        
    conn = get_db()
    cursor = conn.cursor()
    
    query = '''
        SELECT a.*, b.fecha as fecha_origen,
          COALESCE((
            WITH RECURSIVE lineage AS (
              SELECT id, parent_task_id, duracion_min FROM actividades WHERE id = a.id
              UNION ALL
              SELECT prev.id, prev.parent_task_id, prev.duracion_min
              FROM actividades prev
              JOIN lineage l ON prev.id = l.parent_task_id
            )
            SELECT SUM(duracion_min) FROM lineage
          ), a.duracion_min) AS tiempo_acumulado_min
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        WHERE (b.user_id = ? OR b.colaborador = ?)
          AND b.fecha < ?
          AND a.estado != 'completada'
          AND (a.is_deleted IS NULL OR a.is_deleted = 0)
          AND NOT EXISTS (
            SELECT 1 FROM actividades a_child
            JOIN bitacoras b_child ON a_child.bitacora_id = b_child.id
            WHERE a_child.parent_task_id = a.id
              AND (a_child.is_deleted IS NULL OR a_child.is_deleted = 0)
              AND b_child.fecha > b.fecha
          )
        ORDER BY a.id ASC
    '''
    tasks = cursor.execute(query, (user_id, colaborador, target_fecha)).fetchall()
    conn.close()
    
    return jsonify({
        "success": True,
        "pendientes": [format_actividad_dict(t) for t in tasks],
        "count": len(tasks)
    })

# ==================== ACTIVIDADES REFERENCIAS / VÍNCULOS ====================

@app.route('/api/actividades/referencias', methods=['GET', 'OPTIONS'])
def get_actividades_referencias():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    user_id = request.args.get('user_id')
    colaborador = request.args.get('colaborador')
    q = request.args.get('q', '').strip()
    
    conn = get_db()
    cursor = conn.cursor()
    
    query = '''
        SELECT a.id, a.orden, a.hora_inicio, a.duracion_min, a.tipo_trabajo, a.descripcion,
               a.para_cliente, a.estado, b.fecha as bitacora_fecha, b.colaborador
        FROM actividades a
        JOIN bitacoras b ON a.bitacora_id = b.id
        WHERE (a.is_deleted IS NULL OR a.is_deleted = 0)
    '''
    params = []
    if user_id:
        query += " AND (b.user_id = ? OR LOWER(b.colaborador) = LOWER(?))"
        params.extend([user_id, colaborador or ''])
    elif colaborador:
        query += " AND LOWER(b.colaborador) = LOWER(?)"
        params.append(colaborador)

    if q:
        query += " AND (a.descripcion LIKE ? OR a.para_cliente LIKE ? OR a.tipo_trabajo LIKE ?)"
        params.extend([f"%{q}%", f"%{q}%", f"%{q}%"])
        
    query += " ORDER BY b.fecha DESC, a.orden ASC, a.id DESC LIMIT 80"
    rows = cursor.execute(query, params).fetchall()
    conn.close()
    
    return jsonify([dict(r) for r in rows]), 200

# ==================== BUZÓN DE SUGERENCIAS & MEJORA CONTINUA ====================

@app.route('/api/sugerencias', methods=['GET', 'POST', 'OPTIONS'])
def handle_sugerencias():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    now_peru = get_peru_now_str()

    if request.method == 'POST':
        data = request.get_json() or {}
        user_id = data.get('user_id')
        colaborador = (data.get('colaborador') or '').strip()
        es_anonimo = 1 if data.get('es_anonimo') else 0
        categoria = data.get('categoria') or 'sistema'
        titulo = (data.get('titulo') or '').strip()
        descripcion = (data.get('descripcion') or '').strip()
        impacto = data.get('impacto') or 'medio'

        if not titulo or not descripcion:
            conn.close()
            return jsonify({"error": "Título y descripción son requeridos"}), 400

        cursor.execute('''
            INSERT INTO buzon_sugerencias (
                user_id, colaborador, es_anonimo, categoria,
                titulo, descripcion, impacto, estado, votos, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 'pendiente', 0, ?, ?)
        ''', (user_id, colaborador, es_anonimo, categoria, titulo, descripcion, impacto, now_peru, now_peru))
        sug_id = cursor.lastrowid
        conn.commit()

        # Notificar a administradores y líderes
        try:
            author_label = "Un colaborador (Anónimo)" if es_anonimo else (colaborador or "Un colaborador")
            cursor.execute('''
                INSERT INTO notifications (user_id, title, message, type, is_read, created_at)
                SELECT id, 'Nueva sugerencia en el buzón', ?, 'info', 0, ?
                FROM users WHERE role IN ('admin', 'lider')
            ''', (f'{author_label} propuso: "{titulo}"', now_peru))
            conn.commit()
        except Exception as e:
            print("Error creating suggestion notification:", e)

        sug = cursor.execute("SELECT * FROM buzon_sugerencias WHERE id = ?", (sug_id,)).fetchone()
        conn.close()
        return jsonify({"message": "Sugerencia enviada exitosamente", "sugerencia": dict(sug)}), 201

    # GET sugerencias
    categoria = request.args.get('categoria')
    estado = request.args.get('estado')
    user_id = request.args.get('user_id')
    current_uid = request.args.get('requesting_user_id') or user_id

    query = '''
        SELECT s.*,
               CASE WHEN bv.id IS NOT NULL THEN 1 ELSE 0 END as user_has_voted
        FROM buzon_sugerencias s
        LEFT JOIN buzon_votos bv ON s.id = bv.sugerencia_id AND bv.user_id = ?
        WHERE 1=1
    '''
    params = [current_uid]

    if categoria and categoria != 'todas':
        query += " AND s.categoria = ?"
        params.append(categoria)
    if estado and estado != 'todas':
        query += " AND s.estado = ?"
        params.append(estado)
    if request.args.get('mine') == 'true' and user_id:
        query += " AND s.user_id = ?"
        params.append(user_id)

    order_by = request.args.get('sort', 'reciente')
    if order_by == 'popular':
        query += " ORDER BY s.votos DESC, s.created_at DESC"
    else:
        query += " ORDER BY s.created_at DESC, s.id DESC"

    rows = cursor.execute(query, params).fetchall()
    results = []
    for r in rows:
        d = dict(r)
        if d.get('es_anonimo'):
            d['colaborador'] = 'Colaborador Anónimo'
            d['user_id'] = None
        results.append(d)

    conn.close()
    return jsonify(results), 200

@app.route('/api/sugerencias/<int:sug_id>/votar', methods=['POST', 'OPTIONS'])
def votar_sugerencia(sug_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    data = request.get_json() or {}
    user_id = data.get('user_id')
    if not user_id:
        return jsonify({"error": "user_id es requerido para votar"}), 400

    conn = get_db()
    cursor = conn.cursor()

    existing_vote = cursor.execute(
        "SELECT id FROM buzon_votos WHERE sugerencia_id = ? AND user_id = ?",
        (sug_id, user_id)
    ).fetchone()

    if existing_vote:
        cursor.execute("DELETE FROM buzon_votos WHERE id = ?", (existing_vote['id'],))
        cursor.execute("UPDATE buzon_sugerencias SET votos = MAX(0, votos - 1) WHERE id = ?", (sug_id,))
        voted = False
    else:
        cursor.execute("INSERT INTO buzon_votos (sugerencia_id, user_id) VALUES (?, ?)", (sug_id, user_id))
        cursor.execute("UPDATE buzon_sugerencias SET votos = votos + 1 WHERE id = ?", (sug_id,))
        voted = True

    conn.commit()
    sug = cursor.execute("SELECT votos FROM buzon_sugerencias WHERE id = ?", (sug_id,)).fetchone()
    conn.close()

    return jsonify({
        "success": True,
        "voted": voted,
        "votos": sug['votos'] if sug else 0
    }), 200

@app.route('/api/sugerencias/<int:sug_id>/status', methods=['PUT', 'OPTIONS'])
def update_sugerencia_status(sug_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    data = request.get_json() or {}
    nuevo_estado = data.get('estado')
    respuesta_admin = data.get('respuesta_admin')
    respondido_por = data.get('respondido_por') or 'Dirección'
    now_peru = get_peru_now_str()

    if not nuevo_estado:
        return jsonify({"error": "estado es requerido"}), 400

    conn = get_db()
    cursor = conn.cursor()
    sug = cursor.execute("SELECT * FROM buzon_sugerencias WHERE id = ?", (sug_id,)).fetchone()
    if not sug:
        conn.close()
        return jsonify({"error": "Sugerencia no encontrada"}), 404

    cursor.execute('''
        UPDATE buzon_sugerencias SET
            estado = ?,
            respuesta_admin = COALESCE(?, respuesta_admin),
            respondido_por = ?,
            respondido_at = ?,
            updated_at = ?
        WHERE id = ?
    ''', (nuevo_estado, respuesta_admin, respondido_por, now_peru, now_peru, sug_id))
    conn.commit()

    if sug['user_id'] and not sug['es_anonimo']:
        try:
            cursor.execute('''
                INSERT INTO notifications (user_id, title, message, type, is_read, created_at)
                VALUES (?, 'Actualización de tu sugerencia', ?, 'success', 0, ?)
            ''', (sug['user_id'], f'Tu propuesta "{sug["titulo"]}" ahora está: {nuevo_estado.upper()}', now_peru))
            conn.commit()
        except Exception:
            pass

    updated = cursor.execute("SELECT * FROM buzon_sugerencias WHERE id = ?", (sug_id,)).fetchone()
    conn.close()
    return jsonify({"success": True, "sugerencia": dict(updated)}), 200

@app.route('/api/sugerencias/<int:sug_id>', methods=['DELETE', 'OPTIONS'])
def delete_sugerencia(sug_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    data = request.get_json(silent=True) or {}
    user_id = data.get('user_id')
    user_role = data.get('role')

    conn = get_db()
    cursor = conn.cursor()
    sug = cursor.execute("SELECT * FROM buzon_sugerencias WHERE id = ?", (sug_id,)).fetchone()
    if not sug:
        conn.close()
        return jsonify({"error": "Sugerencia no encontrada"}), 404

    if user_role != 'admin' and (not user_id or sug['user_id'] != user_id):
        conn.close()
        return jsonify({"error": "No tienes permiso para eliminar esta sugerencia"}), 403

    cursor.execute("DELETE FROM buzon_votos WHERE sugerencia_id = ?", (sug_id,))
    cursor.execute("DELETE FROM buzon_sugerencias WHERE id = ?", (sug_id,))
    conn.commit()
    return jsonify({"success": True, "message": "Sugerencia eliminada correctamente"}), 200

# ==================== CENTRO DE NOTIFICACIONES ====================

@app.route('/api/notifications', methods=['GET', 'DELETE', 'OPTIONS'])
def handle_notifications():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    user_id = request.args.get('user_id')

    if request.method == 'GET':
        if not user_id:
            conn.close()
            return jsonify([]), 200
        rows = cursor.execute('''
            SELECT id, user_id, title, message, type, is_read, created_at as timestamp
            FROM notifications
            WHERE user_id = ?
            ORDER BY id DESC
            LIMIT 50
        ''', (user_id,)).fetchall()
        result = []
        for r in rows:
            d = dict(r)
            d['id'] = str(d['id'])
            d['read'] = bool(d.pop('is_read', 0))
            result.append(d)
        conn.close()
        return jsonify(result), 200

    elif request.method == 'DELETE':
        if user_id:
            cursor.execute("DELETE FROM notifications WHERE user_id = ?", (user_id,))
            conn.commit()
        conn.close()
        return jsonify({"success": True, "message": "Notificaciones eliminadas"}), 200

@app.route('/api/notifications/<int:notif_id>/read', methods=['PATCH', 'OPTIONS'])
def mark_notification_read(notif_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE notifications SET is_read = 1 WHERE id = ?", (notif_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True}), 200

@app.route('/api/notifications/read-all', methods=['POST', 'OPTIONS'])
def mark_all_notifications_read():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    data = request.get_json(silent=True) or {}
    user_id = data.get('user_id')
    if user_id:
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("UPDATE notifications SET is_read = 1 WHERE user_id = ?", (user_id,))
        conn.commit()
        conn.close()
    return jsonify({"success": True}), 200


# =========================================================================
# RUTAS REST: ACCESOS DIRECTOS (QUICK LINKS)
# =========================================================================

@app.route('/api/quick-links', methods=['GET', 'POST', 'OPTIONS'])
def api_quick_links():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'GET':
        user_id = request.args.get('user_id', type=int)
        categoria = request.args.get('categoria', '').strip()
        search = request.args.get('search', '').strip().lower()

        user_team_id = None
        user_role = 'analista'
        if user_id:
            cursor.execute("SELECT team_id, role FROM users WHERE id = ?", (user_id,))
            u_row = cursor.fetchone()
            if u_row:
                user_team_id = u_row['team_id']
                user_role = u_row['role'] or 'analista'

        query = '''
            SELECT q.*, u.full_name as author_name, t.nombre as team_name
            FROM quick_links q
            LEFT JOIN users u ON q.user_id = u.id
            LEFT JOIN teams t ON q.team_id = t.id
            WHERE 1=1
        '''
        params = []

        if user_role not in ('admin', 'lider'):
            if user_id:
                if user_team_id:
                    query += " AND (q.visibilidad = 'global' OR (q.visibilidad = 'equipo' AND q.team_id = ?) OR (q.visibilidad = 'personal' AND q.user_id = ?))"
                    params.extend([user_team_id, user_id])
                else:
                    query += " AND (q.visibilidad = 'global' OR (q.visibilidad = 'personal' AND q.user_id = ?))"
                    params.append(user_id)
            else:
                query += " AND q.visibilidad = 'global'"

        if categoria and categoria != 'todas':
            query += " AND q.categoria = ?"
            params.append(categoria)

        if search:
            query += " AND (LOWER(q.titulo) LIKE ? OR LOWER(q.url) LIKE ? OR LOWER(COALESCE(q.descripcion, '')) LIKE ?)"
            s_param = f"%{search}%"
            params.extend([s_param, s_param, s_param])

        query += " ORDER BY q.updated_at DESC, q.titulo ASC"
        cursor.execute(query, params)
        raw_links = cursor.fetchall()
        links = []
        for row in raw_links:
            item = dict(row)
            if item.get('password'):
                item['password'] = decrypt_vault_secret(item['password'])
            links.append(item)
        conn.close()
        return jsonify(links), 200

    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        user_id = data.get('user_id')
        titulo = (data.get('titulo') or '').strip()
        url = (data.get('url') or '').strip()

        if not user_id or not titulo or not url:
            conn.close()
            return jsonify({"error": "user_id, titulo y url son obligatorios"}), 400

        cursor.execute("SELECT team_id FROM users WHERE id = ?", (user_id,))
        u_row = cursor.fetchone()
        team_id = u_row['team_id'] if u_row else None

        categoria = (data.get('categoria') or 'General').strip()
        descripcion = (data.get('descripcion') or '').strip()
        icono = (data.get('icono') or 'Link').strip()
        color = (data.get('color') or '#00F0FF').strip()
        usuario = (data.get('usuario') or '').strip()
        raw_password = (data.get('password') or '').strip()
        password = encrypt_vault_secret(raw_password) if raw_password else ''
        visibilidad = data.get('visibilidad', 'personal')
        if visibilidad not in ('personal', 'equipo', 'global'):
            visibilidad = 'personal'

        now_peru = get_peru_now()
        cursor.execute('''
            INSERT INTO quick_links (
                user_id, titulo, url, categoria, descripcion, icono, color,
                usuario, password, visibilidad, team_id, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (user_id, titulo, url, categoria, descripcion, icono, color,
              usuario, password, visibilidad, team_id, now_peru, now_peru))
        new_id = cursor.lastrowid
        conn.commit()

        cursor.execute('''
            SELECT q.*, u.full_name as author_name, t.nombre as team_name
            FROM quick_links q
            LEFT JOIN users u ON q.user_id = u.id
            LEFT JOIN teams t ON q.team_id = t.id
            WHERE q.id = ?
        ''', (new_id,))
        created_item = dict(cursor.fetchone())
        if created_item.get('password'):
            created_item['password'] = decrypt_vault_secret(created_item['password'])
        conn.close()
        return jsonify(created_item), 201


@app.route('/api/quick-links/<int:link_id>', methods=['PUT', 'DELETE', 'OPTIONS'])
def api_quick_link_detail(link_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM quick_links WHERE id = ?", (link_id,))
    link = cursor.fetchone()
    if not link:
        conn.close()
        return jsonify({"error": "Acceso directo no encontrado"}), 404

    if request.method == 'DELETE':
        user_id = request.args.get('user_id', type=int)
        cursor.execute("SELECT role FROM users WHERE id = ?", (user_id,))
        u = cursor.fetchone()
        is_admin = u and u['role'] == 'admin'

        if not is_admin and link['user_id'] != user_id:
            conn.close()
            return jsonify({"error": "No tienes permiso para eliminar este enlace"}), 403

        cursor.execute("DELETE FROM quick_links WHERE id = ?", (link_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True}), 200

    if request.method == 'PUT':
        data = request.get_json(silent=True) or {}
        user_id = data.get('user_id')
        cursor.execute("SELECT role FROM users WHERE id = ?", (user_id,))
        u = cursor.fetchone()
        is_admin = u and u['role'] == 'admin'

        if not is_admin and link['user_id'] != user_id:
            conn.close()
            return jsonify({"error": "No tienes permiso para modificar este enlace"}), 403

        titulo = (data.get('titulo') or link['titulo']).strip()
        url = (data.get('url') or link['url']).strip()
        categoria = (data.get('categoria') or link['categoria']).strip()
        descripcion = (data.get('descripcion') if 'descripcion' in data else link['descripcion']) or ''
        icono = data.get('icono') or link['icono']
        color = data.get('color') or link['color']
        usuario = data.get('usuario') if 'usuario' in data else link['usuario']
        if 'password' in data:
            raw_pass = (data.get('password') or '').strip()
            password = encrypt_vault_secret(raw_pass) if raw_pass else ''
        else:
            password = link['password']
        visibilidad = data.get('visibilidad') or link['visibilidad']

        now_peru = get_peru_now()
        cursor.execute('''
            UPDATE quick_links SET
                titulo = ?, url = ?, categoria = ?, descripcion = ?,
                icono = ?, color = ?, usuario = ?, password = ?,
                visibilidad = ?, updated_at = ?
            WHERE id = ?
        ''', (titulo, url, categoria, descripcion, icono, color, usuario, password, visibilidad, now_peru, link_id))
        conn.commit()

        cursor.execute('''
            SELECT q.*, u.full_name as author_name, t.nombre as team_name
            FROM quick_links q
            LEFT JOIN users u ON q.user_id = u.id
            LEFT JOIN teams t ON q.team_id = t.id
            WHERE q.id = ?
        ''', (link_id,))
        updated_item = dict(cursor.fetchone())
        if updated_item.get('password'):
            updated_item['password'] = decrypt_vault_secret(updated_item['password'])
        conn.close()
        return jsonify(updated_item), 200


# =========================================================================
# RUTAS REST: BÓVEDA IT CORPORATIVA (MULTI-EMPRESA / MARCAS / CREDENCIALES)
# =========================================================================

@app.route('/api/it-vault/empresas', methods=['GET', 'POST', 'OPTIONS'])
def api_it_empresas():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'GET':
        cursor.execute("SELECT * FROM it_empresas ORDER BY nombre ASC")
        empresas = [dict(row) for row in cursor.fetchall()]
        for emp in empresas:
            cursor.execute("SELECT * FROM it_marcas WHERE empresa_id = ? ORDER BY nombre ASC", (emp['id'],))
            emp['marcas'] = [dict(m) for m in cursor.fetchall()]
        conn.close()
        return jsonify(empresas), 200

    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        nombre = (data.get('nombre') or '').strip()
        color = data.get('color', '#00F0FF')
        if not nombre:
            conn.close()
            return jsonify({"error": "El nombre de la empresa es obligatorio"}), 400

        try:
            cursor.execute("INSERT INTO it_empresas (nombre, color) VALUES (?, ?)", (nombre, color))
            new_id = cursor.lastrowid
            conn.commit()
            conn.close()
            return jsonify({"id": new_id, "nombre": nombre, "color": color, "marcas": []}), 201
        except sqlite3.IntegrityError:
            conn.close()
            return jsonify({"error": "Ya existe una empresa con ese nombre"}), 400


@app.route('/api/it-vault/marcas', methods=['POST', 'OPTIONS'])
def api_it_marcas():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    data = request.get_json(silent=True) or {}
    empresa_id = data.get('empresa_id')
    nombre = (data.get('nombre') or '').strip()
    if not empresa_id or not nombre:
        return jsonify({"error": "empresa_id y nombre son obligatorios"}), 400

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO it_marcas (empresa_id, nombre) VALUES (?, ?)", (empresa_id, nombre))
        new_id = cursor.lastrowid
        conn.commit()
        conn.close()
        return jsonify({"id": new_id, "empresa_id": empresa_id, "nombre": nombre}), 201
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "Esta marca ya existe para la empresa seleccionada"}), 400


@app.route('/api/it-vault/credentials', methods=['GET', 'POST', 'OPTIONS'])
def api_it_credentials():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'GET':
        user_id = request.args.get('user_id', type=int)
        empresa_id = request.args.get('empresa_id', type=int)
        marca_id = request.args.get('marca_id', type=int)
        tipo_servicio = request.args.get('tipo_servicio', '').strip()
        search = request.args.get('search', '').strip().lower()

        # Determinar rol del usuario
        user_role = 'analista'
        if user_id:
            cursor.execute("SELECT role FROM users WHERE id = ?", (user_id,))
            ur = cursor.fetchone()
            if ur:
                user_role = ur['role'] or 'analista'

        is_admin_or_leader = user_role in ('admin', 'lider')

        query = '''
            SELECT c.*, 
                   e.nombre as empresa_nombre, e.color as empresa_color,
                   m.nombre as marca_nombre,
                   u.full_name as author_name,
                   COALESCE(p.can_view, 0) as user_can_view,
                   COALESCE(p.can_edit, 0) as user_can_edit
            FROM it_credentials c
            JOIN it_empresas e ON c.empresa_id = e.id
            LEFT JOIN it_marcas m ON c.marca_id = m.id
            LEFT JOIN users u ON c.created_by = u.id
            LEFT JOIN it_credential_permissions p ON p.credential_id = c.id AND p.user_id = ?
            WHERE c.is_active = 1
        '''
        params = [user_id or 0]

        if not is_admin_or_leader:
            query += " AND (c.created_by = ? OR p.can_view = 1)"
            params.append(user_id or 0)

        if empresa_id:
            query += " AND c.empresa_id = ?"
            params.append(empresa_id)

        if marca_id:
            query += " AND c.marca_id = ?"
            params.append(marca_id)

        if tipo_servicio and tipo_servicio != 'todas':
            query += " AND c.tipo_servicio = ?"
            params.append(tipo_servicio)

        if search:
            query += " AND (LOWER(c.plataforma) LIKE ? OR LOWER(c.usuario_login) LIKE ? OR LOWER(COALESCE(c.notas, '')) LIKE ? OR LOWER(e.nombre) LIKE ?)"
            s_param = f"%{search}%"
            params.extend([s_param, s_param, s_param, s_param])

        query += " ORDER BY e.nombre ASC, m.nombre ASC, c.plataforma ASC"
        cursor.execute(query, params)
        rows = cursor.fetchall()
        
        results = []
        for r in rows:
            item = dict(r)
            if item.get('password_secret'):
                item['password_secret'] = decrypt_vault_secret(item['password_secret'])
            # Si es admin o líder, tiene permisos completos
            if is_admin_or_leader or item['created_by'] == user_id:
                item['can_view'] = True
                item['can_edit'] = True
            else:
                item['can_view'] = bool(item.get('user_can_view'))
                item['can_edit'] = bool(item.get('user_can_edit'))
            results.append(item)

        conn.close()
        return jsonify(results), 200

    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        user_id = data.get('created_by') or data.get('user_id')
        empresa_id = data.get('empresa_id')
        marca_id = data.get('marca_id')
        plataforma = (data.get('plataforma') or '').strip()
        tipo_servicio = (data.get('tipo_servicio') or 'General').strip()
        url_acceso = (data.get('url_acceso') or '').strip()
        usuario_login = (data.get('usuario_login') or '').strip()
        raw_secret = (data.get('password_secret') or '').strip()
        password_secret = encrypt_vault_secret(raw_secret) if raw_secret else ''
        notas = (data.get('notas') or '').strip()
        tipo_cuenta = data.get('tipo_cuenta', 'operativa')
        permissions = data.get('permissions', [])  # list of { user_id, can_view, can_edit }

        if not user_id or not empresa_id or not plataforma or not usuario_login or not raw_secret:
            conn.close()
            return jsonify({"error": "empresa_id, plataforma, usuario_login y password_secret son requeridos"}), 400

        now_peru = get_peru_now()
        cursor.execute('''
            INSERT INTO it_credentials (
                empresa_id, marca_id, plataforma, tipo_servicio, url_acceso,
                usuario_login, password_secret, notas, tipo_cuenta, created_by,
                created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (empresa_id, marca_id, plataforma, tipo_servicio, url_acceso,
              usuario_login, password_secret, notas, tipo_cuenta, user_id, now_peru, now_peru))
        cred_id = cursor.lastrowid

        # Permiso completo para el creador
        cursor.execute('''
            INSERT OR REPLACE INTO it_credential_permissions (credential_id, user_id, can_view, can_edit, assigned_by)
            VALUES (?, ?, 1, 1, ?)
        ''', (cred_id, user_id, user_id))

        # Asignar permisos específicos opcionales enviados por el admin
        for p in permissions:
            u_p_id = p.get('user_id')
            if u_p_id and u_p_id != user_id:
                c_view = 1 if p.get('can_view') else 0
                c_edit = 1 if p.get('can_edit') else 0
                cursor.execute('''
                    INSERT OR REPLACE INTO it_credential_permissions (credential_id, user_id, can_view, can_edit, assigned_by)
                    VALUES (?, ?, ?, ?, ?)
                ''', (cred_id, u_p_id, c_view, c_edit, user_id))

        # Auditoría de creación
        cursor.execute('''
            INSERT INTO it_credential_audit (credential_id, user_id, action, ip_or_agent)
            VALUES (?, ?, 'create', ?)
        ''', (cred_id, user_id, request.remote_addr))

        conn.commit()
        conn.close()
        return jsonify({"id": cred_id, "success": True}), 201


@app.route('/api/it-vault/credentials/<int:cred_id>', methods=['PUT', 'DELETE', 'OPTIONS'])
def api_it_credential_detail(cred_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM it_credentials WHERE id = ? AND is_active = 1", (cred_id,))
    cred = cursor.fetchone()
    if not cred:
        conn.close()
        return jsonify({"error": "Credencial no encontrada"}), 404

    if request.method == 'DELETE':
        user_id = request.args.get('user_id', type=int)
        cursor.execute("SELECT role FROM users WHERE id = ?", (user_id,))
        u = cursor.fetchone()
        is_admin = u and u['role'] in ('admin', 'lider')

        if not is_admin and cred['created_by'] != user_id:
            conn.close()
            return jsonify({"error": "No tienes permiso para eliminar esta credencial"}), 403

        cursor.execute("UPDATE it_credentials SET is_active = 0 WHERE id = ?", (cred_id,))
        cursor.execute('''
            INSERT INTO it_credential_audit (credential_id, user_id, action, ip_or_agent)
            VALUES (?, ?, 'delete', ?)
        ''', (cred_id, user_id, request.remote_addr))
        conn.commit()
        conn.close()
        return jsonify({"success": True}), 200

    if request.method == 'PUT':
        data = request.get_json(silent=True) or {}
        user_id = data.get('updated_by') or data.get('user_id')
        cursor.execute("SELECT role FROM users WHERE id = ?", (user_id,))
        u = cursor.fetchone()
        is_admin = u and u['role'] in ('admin', 'lider')

        # Verificar permiso can_edit
        cursor.execute("SELECT can_edit FROM it_credential_permissions WHERE credential_id = ? AND user_id = ?", (cred_id, user_id))
        perm = cursor.fetchone()
        can_edit = (perm and perm['can_edit'] == 1) or is_admin or (cred['created_by'] == user_id)

        if not can_edit:
            conn.close()
            return jsonify({"error": "No tienes permiso de edición sobre esta credencial"}), 403

        empresa_id = data.get('empresa_id') or cred['empresa_id']
        marca_id = data.get('marca_id') if 'marca_id' in data else cred['marca_id']
        plataforma = (data.get('plataforma') or cred['plataforma']).strip()
        tipo_servicio = (data.get('tipo_servicio') or cred['tipo_servicio']).strip()
        url_acceso = data.get('url_acceso') if 'url_acceso' in data else cred['url_acceso']
        usuario_login = (data.get('usuario_login') or cred['usuario_login']).strip()
        if 'password_secret' in data:
            raw_secret = (data.get('password_secret') or '').strip()
            password_secret = encrypt_vault_secret(raw_secret) if raw_secret else ''
        else:
            password_secret = cred['password_secret']
        notas = data.get('notas') if 'notas' in data else cred['notas']
        tipo_cuenta = data.get('tipo_cuenta') or cred['tipo_cuenta']

        now_peru = get_peru_now()
        cursor.execute('''
            UPDATE it_credentials SET
                empresa_id = ?, marca_id = ?, plataforma = ?, tipo_servicio = ?,
                url_acceso = ?, usuario_login = ?, password_secret = ?,
                notas = ?, tipo_cuenta = ?, updated_by = ?, updated_at = ?
            WHERE id = ?
        ''', (empresa_id, marca_id, plataforma, tipo_servicio, url_acceso,
              usuario_login, password_secret, notas, tipo_cuenta, user_id, now_peru, cred_id))

        cursor.execute('''
            INSERT INTO it_credential_audit (credential_id, user_id, action, ip_or_agent)
            VALUES (?, ?, 'update', ?)
        ''', (cred_id, user_id, request.remote_addr))

        conn.commit()
        conn.close()
        return jsonify({"success": True}), 200


@app.route('/api/it-vault/credentials/<int:cred_id>/audit', methods=['POST', 'OPTIONS'])
def api_it_credential_audit(cred_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    data = request.get_json(silent=True) or {}
    user_id = data.get('user_id')
    action = data.get('action') or 'view_password'
    if not user_id:
        return jsonify({"error": "user_id es requerido"}), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO it_credential_audit (credential_id, user_id, action, ip_or_agent)
        VALUES (?, ?, ?, ?)
    ''', (cred_id, user_id, action, request.remote_addr))
    conn.commit()
    conn.close()
    return jsonify({"success": True}), 201


@app.route('/api/it-vault/credentials/<int:cred_id>/permissions', methods=['GET', 'PUT', 'OPTIONS'])
def api_it_credential_permissions(cred_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'GET':
        cursor.execute('''
            SELECT u.id as user_id, u.full_name, u.username, u.role,
                   COALESCE(p.can_view, 0) as can_view,
                   COALESCE(p.can_edit, 0) as can_edit
            FROM users u
            LEFT JOIN it_credential_permissions p ON p.user_id = u.id AND p.credential_id = ?
            WHERE u.is_active = 1
            ORDER BY u.full_name ASC
        ''', (cred_id,))
        permissions = [dict(row) for row in cursor.fetchall()]
        conn.close()
        return jsonify(permissions), 200

    if request.method == 'PUT':
        data = request.get_json(silent=True) or {}
        assigned_by = data.get('assigned_by')
        permissions_list = data.get('permissions', [])

        cursor.execute("SELECT role FROM users WHERE id = ?", (assigned_by,))
        u = cursor.fetchone()
        if not u or u['role'] not in ('admin', 'lider'):
            conn.close()
            return jsonify({"error": "Solo administradores o líderes pueden configurar permisos"}), 403

        for perm in permissions_list:
            target_user_id = perm.get('user_id')
            can_view = 1 if perm.get('can_view') else 0
            can_edit = 1 if perm.get('can_edit') else 0
            cursor.execute('''
                INSERT INTO it_credential_permissions (credential_id, user_id, can_view, can_edit, assigned_by)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(credential_id, user_id) DO UPDATE SET
                    can_view = excluded.can_view,
                    can_edit = excluded.can_edit,
                    assigned_by = excluded.assigned_by
            ''', (cred_id, target_user_id, can_view, can_edit, assigned_by))

        conn.commit()
        conn.close()
        return jsonify({"success": True}), 200


# =========================================================================
# RUTAS REST: DIRECTORIO DE USUARIOS POR PLATAFORMA (SOPORTE Y RESETEO)
# =========================================================================

@app.route('/api/it-vault/platform-users', methods=['GET', 'POST', 'OPTIONS'])
def api_it_platform_users():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'GET':
        empresa_id = request.args.get('empresa_id', type=int)
        marca_id = request.args.get('marca_id', type=int)
        plataforma = request.args.get('plataforma', '').strip()
        search = request.args.get('search', '').strip().lower()
        estado = request.args.get('estado', '').strip()

        query = '''
            SELECT pu.*, 
                   e.nombre as empresa_nombre, e.color as empresa_color,
                   m.nombre as marca_nombre,
                   u.full_name as created_by_name
            FROM it_platform_users pu
            JOIN it_empresas e ON pu.empresa_id = e.id
            LEFT JOIN it_marcas m ON pu.marca_id = m.id
            LEFT JOIN users u ON pu.created_by = u.id
            WHERE 1=1
        '''
        params = []

        if empresa_id:
            query += " AND pu.empresa_id = ?"
            params.append(empresa_id)

        if marca_id:
            query += " AND pu.marca_id = ?"
            params.append(marca_id)

        if plataforma and plataforma != 'todas':
            query += " AND pu.plataforma = ?"
            params.append(plataforma)

        if estado and estado != 'todos':
            query += " AND pu.estado = ?"
            params.append(estado)

        if search:
            query += " AND (LOWER(pu.colaborador_nombre) LIKE ? OR LOWER(pu.usuario_login) LIKE ? OR LOWER(COALESCE(pu.colaborador_email, '')) LIKE ? OR LOWER(COALESCE(pu.colaborador_cargo, '')) LIKE ?)"
            s_param = f"%{search}%"
            params.extend([s_param, s_param, s_param, s_param])

        query += " ORDER BY pu.colaborador_nombre ASC, pu.plataforma ASC"
        cursor.execute(query, params)
        raw_users = cursor.fetchall()
        users_list = []
        for row in raw_users:
            item = dict(row)
            if item.get('password_actual'):
                item['password_actual'] = decrypt_vault_secret(item['password_actual'])
            if item.get('password_anterior'):
                item['password_anterior'] = decrypt_vault_secret(item['password_anterior'])
            users_list.append(item)
        conn.close()
        return jsonify(users_list), 200

    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        empresa_id = data.get('empresa_id')
        marca_id = data.get('marca_id') or None
        plataforma = (data.get('plataforma') or '').strip()
        colaborador_nombre = (data.get('colaborador_nombre') or '').strip()
        colaborador_cargo = (data.get('colaborador_cargo') or '').strip()
        colaborador_email = (data.get('colaborador_email') or '').strip()
        colaborador_telefono = (data.get('colaborador_telefono') or '').strip()
        usuario_login = (data.get('usuario_login') or '').strip()
        raw_password = (data.get('password_actual') or '').strip()
        password_actual = encrypt_vault_secret(raw_password) if raw_password else ''
        estado = data.get('estado') or 'activo'
        notas = (data.get('notas') or '').strip()
        created_by = data.get('created_by') or 1

        if not empresa_id or not plataforma or not colaborador_nombre or not usuario_login or not raw_password:
            conn.close()
            return jsonify({"error": "Empresa, Plataforma, Colaborador, Usuario y Contraseña son obligatorios"}), 400

        now_peru = get_peru_now()
        cursor.execute('''
            INSERT INTO it_platform_users (
                empresa_id, marca_id, plataforma, colaborador_nombre, colaborador_cargo,
                colaborador_email, colaborador_telefono, usuario_login, password_actual,
                estado, ultimo_reseteo, notas, created_by, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (empresa_id, marca_id, plataforma, colaborador_nombre, colaborador_cargo,
              colaborador_email, colaborador_telefono, usuario_login, password_actual,
              estado, now_peru, notas, created_by, now_peru, now_peru))
        new_id = cursor.lastrowid
        conn.commit()

        cursor.execute('''
            SELECT pu.*, e.nombre as empresa_nombre, e.color as empresa_color, m.nombre as marca_nombre
            FROM it_platform_users pu
            JOIN it_empresas e ON pu.empresa_id = e.id
            LEFT JOIN it_marcas m ON pu.marca_id = m.id
            WHERE pu.id = ?
        ''', (new_id,))
        created_row = dict(cursor.fetchone())
        if created_row.get('password_actual'):
            created_row['password_actual'] = decrypt_vault_secret(created_row['password_actual'])
        conn.close()
        return jsonify(created_row), 201


@app.route('/api/it-vault/platform-users/<int:user_id>', methods=['PUT', 'DELETE', 'OPTIONS'])
def api_it_platform_user_detail(user_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM it_platform_users WHERE id = ?", (user_id,))
    existing = cursor.fetchone()
    if not existing:
        conn.close()
        return jsonify({"error": "Cuenta de usuario no encontrada"}), 404

    if request.method == 'DELETE':
        cursor.execute("DELETE FROM it_platform_users WHERE id = ?", (user_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True}), 200

    if request.method == 'PUT':
        data = request.get_json(silent=True) or {}
        empresa_id = data.get('empresa_id') or existing['empresa_id']
        marca_id = data.get('marca_id') if 'marca_id' in data else existing['marca_id']
        plataforma = (data.get('plataforma') or existing['plataforma']).strip()
        colaborador_nombre = (data.get('colaborador_nombre') or existing['colaborador_nombre']).strip()
        colaborador_cargo = (data.get('colaborador_cargo') if 'colaborador_cargo' in data else existing['colaborador_cargo']) or ''
        colaborador_email = (data.get('colaborador_email') if 'colaborador_email' in data else existing['colaborador_email']) or ''
        colaborador_telefono = (data.get('colaborador_telefono') if 'colaborador_telefono' in data else existing['colaborador_telefono']) or ''
        usuario_login = (data.get('usuario_login') or existing['usuario_login']).strip()
        if 'password_actual' in data:
            raw_password = (data.get('password_actual') or '').strip()
            password_actual = encrypt_vault_secret(raw_password) if raw_password else ''
        else:
            password_actual = existing['password_actual']
        estado = data.get('estado') or existing['estado']
        notas = (data.get('notas') if 'notas' in data else existing['notas']) or ''
        updated_by = data.get('updated_by')

        now_peru = get_peru_now()
        cursor.execute('''
            UPDATE it_platform_users SET
                empresa_id = ?, marca_id = ?, plataforma = ?, colaborador_nombre = ?,
                colaborador_cargo = ?, colaborador_email = ?, colaborador_telefono = ?,
                usuario_login = ?, password_actual = ?, estado = ?, notas = ?,
                updated_by = ?, updated_at = ?
            WHERE id = ?
        ''', (empresa_id, marca_id, plataforma, colaborador_nombre, colaborador_cargo,
              colaborador_email, colaborador_telefono, usuario_login, password_actual,
              estado, notas, updated_by, now_peru, user_id))
        conn.commit()

        cursor.execute('''
            SELECT pu.*, e.nombre as empresa_nombre, e.color as empresa_color, m.nombre as marca_nombre
            FROM it_platform_users pu
            JOIN it_empresas e ON pu.empresa_id = e.id
            LEFT JOIN it_marcas m ON pu.marca_id = m.id
            WHERE pu.id = ?
        ''', (user_id,))
        updated_row = dict(cursor.fetchone())
        if updated_row.get('password_actual'):
            updated_row['password_actual'] = decrypt_vault_secret(updated_row['password_actual'])
        if updated_row.get('password_anterior'):
            updated_row['password_anterior'] = decrypt_vault_secret(updated_row['password_anterior'])
        conn.close()
        return jsonify(updated_row), 200


@app.route('/api/it-vault/platform-users/<int:user_id>/reset-password', methods=['POST', 'OPTIONS'])
def api_it_platform_user_reset_password(user_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT pu.*, e.nombre as empresa_nombre FROM it_platform_users pu JOIN it_empresas e ON pu.empresa_id = e.id WHERE pu.id = ?", (user_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Cuenta de usuario no encontrada"}), 404

    data = request.get_json(silent=True) or {}
    new_password = (data.get('new_password') or '').strip()
    reset_by = data.get('reset_by')

    # Si no se envió clave, generamos una clave segura automáticamente
    if not new_password:
        import secrets
        import string
        alphabet = string.ascii_letters + string.digits + "!@#$%&*"
        new_password = ''.join(secrets.choice(alphabet) for _ in range(14))

    now_peru = get_peru_now()
    encrypted_new = encrypt_vault_secret(new_password)
    cursor.execute('''
        UPDATE it_platform_users SET
            password_anterior = password_actual,
            password_actual = ?,
            ultimo_reseteo = ?,
            updated_by = ?,
            updated_at = ?
        WHERE id = ?
    ''', (encrypted_new, now_peru, reset_by, now_peru, user_id))
    conn.commit()

    # Preparar plantilla de mensaje de soporte para compartir al colaborador por WhatsApp o correo
    colab = row['colaborador_nombre']
    plat = row['plataforma']
    login = row['usuario_login']
    emp = row['empresa_nombre']

    share_message = (
        f"Hola {colab}, el área de Sistemas de {emp} te comparte tus credenciales actualizadas:\n\n"
        f"Plataforma: {plat}\n"
        f"Usuario: {login}\n"
        f"Contraseña temporal: {new_password}\n\n"
        f"Por seguridad, te sugerimos cambiar tu contraseña al iniciar sesión o mantenerla en un lugar seguro."
    )

    conn.close()
    return jsonify({
        "success": True,
        "new_password": new_password,
        "ultimo_reseteo": now_peru,
        "share_message": share_message
    }), 200


# =========================================================================
# EXCEL & IMPORTACIÓN MASIVA: CUENTAS DE USUARIOS POR PLATAFORMA
# =========================================================================

def normalize_excel_header(h):
    if not h:
        return ""
    h = str(h).strip().lower()
    for char, rep in [("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n"), ("*", ""), (" ", "_"), ("-", "_"), (".", "")]:
        h = h.replace(char, rep)
    return h.strip('_')


def map_row_to_account_dict(raw_dict):
    res = {
        'empresa': '',
        'marca': '',
        'plataforma': '',
        'colaborador_nombre': '',
        'colaborador_cargo': '',
        'colaborador_email': '',
        'colaborador_telefono': '',
        'usuario_login': '',
        'password_actual': '',
        'estado': 'activo',
        'notas': '',
    }
    
    for key, val in raw_dict.items():
        k = normalize_excel_header(key)
        v = str(val).strip() if val is not None else ""
        if not v or v.lower() == 'none':
            v = ""
            
        if k in ['empresa', 'empresa_nombre', 'company']:
            res['empresa'] = v
        elif k in ['marca', 'marca_nombre', 'brand']:
            res['marca'] = v
        elif k in ['plataforma', 'platform', 'servicio', 'sistema']:
            res['plataforma'] = v
        elif k in ['colaborador', 'colaborador_nombre', 'nombre', 'nombre_colaborador', 'empleado', 'usuario_nombre']:
            res['colaborador_nombre'] = v
        elif k in ['cargo', 'colaborador_cargo', 'puesto', 'rol', 'posicion']:
            res['colaborador_cargo'] = v
        elif k in ['email', 'email_colaborador', 'correo', 'correo_colaborador', 'colaborador_email', 'mail']:
            res['colaborador_email'] = v
        elif k in ['telefono', 'colaborador_telefono', 'celular', 'whatsapp', 'phone']:
            res['colaborador_telefono'] = v
        elif k in ['usuario_login', 'usuario', 'login', 'user', 'cuenta', 'username']:
            res['usuario_login'] = v
        elif k in ['password', 'password_actual', 'contrasena', 'clave', 'pass', 'clave_acceso']:
            res['password_actual'] = v
        elif k in ['estado', 'status']:
            st = v.lower()
            if any(x in st for x in ['crear', 'alta', 'pendiente', 'nueva']):
                res['estado'] = 'por_crear'
            elif any(x in st for x in ['susp', 'paus', 'inact']):
                res['estado'] = 'suspendido'
            elif any(x in st for x in ['baja', 'elim', 'cesad']):
                res['estado'] = 'baja'
            else:
                res['estado'] = 'activo'
        elif k in ['notas', 'nota', 'observaciones', 'observacion', 'comentarios']:
            res['notas'] = v

    return res


def parse_accounts_from_filestorage(file_storage):
    filename = file_storage.filename or ''
    lower_fn = filename.lower()
    raw_dicts = []
    
    if lower_fn.endswith('.csv'):
        raw_bytes = file_storage.read()
        decoded = None
        for enc in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']:
            try:
                decoded = raw_bytes.decode(enc)
                break
            except Exception:
                pass
        if decoded is None:
            raise ValueError("No se pudo decodificar el archivo CSV.")
        
        sample = decoded[:2048]
        delimiter = ';' if sample.count(';') > sample.count(',') else ','
        reader = csv.reader(io.StringIO(decoded), delimiter=delimiter)
        rows = list(reader)
        if not rows:
            return []
            
        header_idx = -1
        for i, r in enumerate(rows[:20]):
            non_empty = [c for c in r if c is not None and str(c).strip()]
            if len(non_empty) >= 3:
                r_str = " ".join([str(c).lower() for c in non_empty])
                if 'empresa' in r_str and ('colaborador' in r_str or 'usuario' in r_str or 'plataforma' in r_str):
                    header_idx = i
                    break
        if header_idx == -1:
            for i, r in enumerate(rows[:20]):
                if len([c for c in r if c is not None and str(c).strip()]) >= 3:
                    header_idx = i
                    break
        if header_idx == -1:
            header_idx = 0
                
        headers = [normalize_excel_header(c) for c in rows[header_idx]]
        for r in rows[header_idx + 1:]:
            if not any(str(c).strip() for c in r):
                continue
            item = {}
            for idx, h in enumerate(headers):
                if h and idx < len(r):
                    item[h] = str(r[idx]).strip()
            if item:
                raw_dicts.append(item)
    else:
        if not _HAS_OPENPYXL:
            raise ValueError("El servidor requiere 'openpyxl' para procesar archivos .xlsx. Por favor suba su archivo en formato .csv delimitado por comas o punto y coma, o ejecute 'pip install openpyxl' en el servidor.")
        wb = openpyxl.load_workbook(file_storage, data_only=True)
        ws = None
        for name in ['Cuentas_Plataforma', 'Cuentas', 'Usuarios', 'Hoja1', 'Sheet1']:
            if name in wb.sheetnames:
                ws = wb[name]
                break
        if ws is None:
            ws = wb.active
            
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
            
        header_idx = -1
        for i, r in enumerate(rows[:20]):
            non_empty = [c for c in r if c is not None and str(c).strip()]
            if len(non_empty) >= 3:
                r_str = " ".join([str(c or '').lower() for c in non_empty])
                if 'empresa' in r_str and ('colaborador' in r_str or 'usuario' in r_str or 'plataforma' in r_str):
                    header_idx = i
                    break
        if header_idx == -1:
            for i, r in enumerate(rows[:20]):
                if len([c for c in r if c is not None and str(c).strip()]) >= 3:
                    header_idx = i
                    break
        if header_idx == -1:
            header_idx = 0
                
        headers = [normalize_excel_header(c) for c in rows[header_idx]]
        for r in rows[header_idx + 1:]:
            if not any(c is not None and str(c).strip() for c in r):
                continue
            item = {}
            for idx, h in enumerate(headers):
                if h and idx < len(r):
                    val = r[idx]
                    item[h] = str(val).strip() if val is not None else ""
            if item:
                raw_dicts.append(item)
                
    mapped = []
    for r in raw_dicts:
        acc = map_row_to_account_dict(r)
        if not acc['empresa'] and not acc['plataforma'] and not acc['colaborador_nombre'] and not acc['usuario_login']:
            continue
        mapped.append(acc)
    return mapped


@app.route('/api/it-vault/platform-users/template', methods=['GET'])
def api_it_platform_users_template():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT nombre FROM it_empresas ORDER BY nombre ASC")
    db_empresas = [r[0] for r in cursor.fetchall()]
    cursor.execute("SELECT nombre, tipo_servicio FROM it_plataformas ORDER BY nombre ASC")
    db_plataformas = cursor.fetchall()
    conn.close()

    if not _HAS_OPENPYXL:
        output = io.StringIO()
        output.write('\ufeff')
        writer = csv.writer(output, delimiter=';')
        writer.writerow(["Empresa *", "Marca", "Plataforma *", "Colaborador *", "Cargo", "Email_Colaborador", "Telefono", "Usuario_Login *", "Password *", "Estado", "Notas"])
        def_emp = db_empresas[0] if db_empresas else "Marketing Alterno"
        writer.writerow([def_emp, def_emp, "Zimbra Mail", "Juan Carlos Pérez Flores", "Diseñador Gráfico", "juan.perez@alterno.pe", "+51 987654321", "jperez@alterno.pe", "Temp#Pass2026!", "activo", "Buzón de correo"])
        writer.writerow([def_emp, def_emp, "Odoo ERP", "María Elena Gómez Silva", "Analista Contable", "mgomez@alterno.pe", "+51 912345678", "mgomez", "Odoo#Mkt2026!", "activo", "Módulos de compras"])
        buf = io.BytesIO(output.getvalue().encode('utf-8-sig'))
        return send_file(buf, mimetype='text/csv', as_attachment=True, download_name='Plantilla_Cuentas_Plataforma.csv')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cuentas_Plataforma"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título y Banner
    ws.merge_cells('A1:K1')
    ws['A1'] = "SISTEMAS TI - PLANTILLA DE IMPORTACIÓN MASIVA DE CUENTAS POR PLATAFORMA"
    ws['A1'].font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    ws.merge_cells('A2:K2')
    ws['A2'] = "Instrucciones: Complete los datos a partir de la fila 4. Las columnas con (*) son obligatorias. Las contraseñas serán cifradas automáticamente con AES-256."
    ws['A2'].font = Font(name='Calibri', size=10, italic=True, color='475569')
    ws['A2'].fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[2].height = 24

    headers = [
        ("Empresa *", 20),
        ("Marca", 18),
        ("Plataforma *", 24),
        ("Colaborador *", 28),
        ("Cargo", 24),
        ("Email_Colaborador", 26),
        ("Telefono", 18),
        ("Usuario_Login *", 26),
        ("Password *", 20),
        ("Estado", 16),
        ("Notas", 32)
    ]

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.row_dimensions[3].height = 28
    for col_idx, (h_title, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # Filas de Ejemplo
    def_emp = db_empresas[0] if db_empresas else "Marketing Alterno"
    examples = [
        (def_emp, def_emp, "Zimbra Mail", "Juan Carlos Pérez Flores", "Diseñador Gráfico Senior", "juan.perez@alterno.pe", "+51 987654321", "jperez@alterno.pe", "Temp#Pass2026!", "activo", "Buzón de correo corporativo"),
        (def_emp, def_emp, "Odoo ERP", "María Elena Gómez Silva", "Analista Contable", "mgomez@alterno.pe", "+51 912345678", "mgomez", "Odoo#Mkt2026!", "activo", "Acceso a módulos de compras y facturación"),
        (def_emp, "", "Office 365", "Carlos Alberto Mendoza", "Asistente de Operaciones", "cmendoza@alterno.pe", "", "cmendoza@alterno.pe", "Init#2026M365", "por_crear", "Cuenta en proceso de alta técnica"),
    ]

    example_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    for row_idx, row_data in enumerate(examples, start=4):
        ws.row_dimensions[row_idx].height = 22
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10, color='1E293B')
            cell.fill = example_fill
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal='center' if col_idx in [1, 2, 7, 10] else 'left',
                vertical='center'
            )

    # Hoja 2: Catálogos y Guía
    ws2 = wb.create_sheet(title="Catalogos_Y_Guia")
    ws2.views.sheetView[0].showGridLines = True
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 36
    ws2.column_dimensions['C'].width = 36

    ws2['A1'] = "CATÁLOGOS ACTUALES & GUÍA DE CARGA"
    ws2['A1'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    ws2['A1'].fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    ws2.row_dimensions[1].height = 28

    ws2['A3'] = "EMPRESAS EN BASE DE DATOS"
    ws2['A3'].font = Font(name='Calibri', size=11, bold=True, color='1E293B')
    for idx, emp_name in enumerate(db_empresas, start=4):
        ws2[f'A{idx}'] = emp_name

    plat_start = len(db_empresas) + 5
    ws2[f'A{plat_start}'] = "PLATAFORMAS REGISTRADAS"
    ws2[f'B{plat_start}'] = "TIPO DE SERVICIO"
    ws2[f'A{plat_start}'].font = Font(name='Calibri', size=11, bold=True, color='1E293B')
    ws2[f'B{plat_start}'].font = Font(name='Calibri', size=11, bold=True, color='1E293B')

    for idx, p_row in enumerate(db_plataformas, start=plat_start + 1):
        ws2[f'A{idx}'] = p_row[0]
        ws2[f'B{idx}'] = p_row[1]

    est_start = plat_start + len(db_plataformas) + 3
    ws2[f'A{est_start}'] = "ESTADOS VÁLIDOS"
    ws2[f'B{est_start}'] = "DESCRIPCIÓN"
    ws2[f'A{est_start}'].font = Font(name='Calibri', size=11, bold=True, color='1E293B')
    ws2[f'B{est_start}'].font = Font(name='Calibri', size=11, bold=True, color='1E293B')

    valid_estados = [
        ("activo", "Cuenta operativa y en uso activo por el colaborador."),
        ("por_crear", "Cuenta planificada o pendiente de aprovisionamiento en la plataforma."),
        ("suspendido", "Cuenta bloqueada o pausada por vacaciones o licencia temporal."),
        ("baja", "Cuenta deshabilitada por cese o desvinculación definitiva.")
    ]
    for idx, (st_name, st_desc) in enumerate(valid_estados, start=est_start + 1):
        ws2[f'A{idx}'] = st_name
        ws2[f'B{idx}'] = st_desc

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='Plantilla_Cuentas_Plataforma.xlsx'
    )


@app.route('/api/it-vault/platform-users/export', methods=['GET'])
def api_it_platform_users_export():
    empresa_id = request.args.get('empresa_id', type=int)
    marca_id = request.args.get('marca_id', type=int)
    plataforma = request.args.get('plataforma', '').strip()
    search = request.args.get('search', '').strip().lower()
    estado = request.args.get('estado', '').strip()
    include_passwords = request.args.get('include_passwords', '0') == '1'

    conn = get_db()
    cursor = conn.cursor()

    query = '''
        SELECT pu.*, 
               e.nombre as empresa_nombre,
               m.nombre as marca_nombre,
               u.full_name as created_by_name
        FROM it_platform_users pu
        JOIN it_empresas e ON pu.empresa_id = e.id
        LEFT JOIN it_marcas m ON pu.marca_id = m.id
        LEFT JOIN users u ON pu.created_by = u.id
        WHERE 1=1
    '''
    params = []

    if empresa_id:
        query += " AND pu.empresa_id = ?"
        params.append(empresa_id)

    if marca_id:
        query += " AND pu.marca_id = ?"
        params.append(marca_id)

    if plataforma and plataforma != 'todas':
        query += " AND pu.plataforma = ?"
        params.append(plataforma)

    if estado and estado != 'todos':
        query += " AND pu.estado = ?"
        params.append(estado)

    if search:
        query += " AND (LOWER(pu.colaborador_nombre) LIKE ? OR LOWER(pu.usuario_login) LIKE ? OR LOWER(COALESCE(pu.colaborador_email, '')) LIKE ? OR LOWER(COALESCE(pu.colaborador_cargo, '')) LIKE ?)"
        s_param = f"%{search}%"
        params.extend([s_param, s_param, s_param, s_param])

    query += " ORDER BY e.nombre ASC, pu.plataforma ASC, pu.colaborador_nombre ASC"
    cursor.execute(query, params)
    raw_users = cursor.fetchall()
    conn.close()

    if not _HAS_OPENPYXL:
        output = io.StringIO()
        output.write('\ufeff')
        writer = csv.writer(output, delimiter=';')
        writer.writerow(["ID", "Empresa", "Marca", "Plataforma", "Colaborador", "Cargo", "Email", "Teléfono", "Usuario Login", "Contraseña", "Estado", "Notas"])
        for r in raw_users:
            pw_val = "••••••••"
            if include_passwords and r['password_actual']:
                try:
                    pw_val = decrypt_vault_secret(r['password_actual'])
                except Exception:
                    pw_val = "••••••••"
            writer.writerow([r['id'], r['empresa_nombre'] or '', r['marca_nombre'] or '', r['plataforma'] or '', r['colaborador_nombre'] or '', r['colaborador_cargo'] or '', r['colaborador_email'] or '', r['colaborador_telefono'] or '', r['usuario_login'] or '', pw_val, r['estado'] or 'activo', r['notas'] or ''])
        buf = io.BytesIO(output.getvalue().encode('utf-8-sig'))
        return send_file(buf, mimetype='text/csv', as_attachment=True, download_name='Cuentas_Plataformas_Export.csv')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cuentas_Export"
    ws.views.sheetView[0].showGridLines = True

    # Banner
    ws.merge_cells('A1:L1')
    ws['A1'] = "DIRECTORIO DE CUENTAS POR PLATAFORMA - SISTEMAS / TI"
    ws['A1'].font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    now_str = get_peru_now_str()
    ws.merge_cells('A2:L2')
    ws['A2'] = f"Reporte generado el {now_str} | Total de registros exportados: {len(raw_users)}"
    ws['A2'].font = Font(name='Calibri', size=10, italic=True, color='475569')
    ws['A2'].fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[2].height = 24

    cols = [
        ("ID", 8),
        ("Empresa", 20),
        ("Marca", 16),
        ("Plataforma", 22),
        ("Colaborador", 28),
        ("Cargo", 22),
        ("Email", 26),
        ("Teléfono", 16),
        ("Usuario Login", 24),
        ("Contraseña", 18),
        ("Estado", 14),
        ("Notas", 30)
    ]

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.row_dimensions[3].height = 26
    for col_idx, (col_name, col_width) in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    for row_idx, r in enumerate(raw_users, start=4):
        ws.row_dimensions[row_idx].height = 20
        pw_val = "••••••••"
        if include_passwords and r['password_actual']:
            try:
                pw_val = decrypt_vault_secret(r['password_actual'])
            except Exception:
                pw_val = "••••••••"

        row_vals = [
            r['id'],
            r['empresa_nombre'] or '',
            r['marca_nombre'] or '',
            r['plataforma'] or '',
            r['colaborador_nombre'] or '',
            r['colaborador_cargo'] or '',
            r['colaborador_email'] or '',
            r['colaborador_telefono'] or '',
            r['usuario_login'] or '',
            pw_val,
            r['estado'] or 'activo',
            r['notas'] or ''
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10, color='1E293B')
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal='center' if col_idx in [1, 8, 10, 11] else 'left',
                vertical='center'
            )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='Cuentas_Plataformas_Export.xlsx'
    )


@app.route('/api/it-vault/platform-users/import/preview', methods=['POST', 'OPTIONS'])
def api_it_platform_users_import_preview():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    if 'file' not in request.files:
        return jsonify({"error": "No se envió ningún archivo para previsualización"}), 400

    uploaded_file = request.files['file']
    if not uploaded_file.filename:
        return jsonify({"error": "Archivo no seleccionado"}), 400

    try:
        parsed_rows = parse_accounts_from_filestorage(uploaded_file)
    except Exception as e:
        return jsonify({"error": f"Error al procesar el archivo: {str(e)}"}), 400

    if not parsed_rows:
        return jsonify({"error": "No se encontraron filas con datos en el archivo subido"}), 400

    conn = get_db()
    cursor = conn.cursor()

    # Pre-cargar empresas y cuentas para verificación rápida
    cursor.execute("SELECT id, LOWER(nombre) FROM it_empresas")
    existing_empresas = {r[1]: r[0] for r in cursor.fetchall()}

    cursor.execute("SELECT empresa_id, LOWER(plataforma), LOWER(usuario_login), id, colaborador_nombre FROM it_platform_users")
    existing_accounts = {}
    for r in cursor.fetchall():
        key = (r[0], r[1].strip(), r[2].strip())
        existing_accounts[key] = {"id": r[3], "colaborador_nombre": r[4]}

    conn.close()

    preview_rows = []
    ready_count = 0
    update_count = 0
    invalid_count = 0

    for idx, r in enumerate(parsed_rows, start=1):
        errs = []
        if not r.get('empresa'):
            errs.append("Empresa es obligatoria")
        if not r.get('plataforma'):
            errs.append("Plataforma es obligatoria")
        if not r.get('colaborador_nombre'):
            errs.append("Colaborador es obligatorio")
        if not r.get('usuario_login'):
            errs.append("Usuario Login es obligatorio")

        emp_lower = (r.get('empresa') or '').strip().lower()
        emp_id = existing_empresas.get(emp_lower)

        status = 'ready'
        status_message = 'Listo para registrar'

        if errs:
            status = 'invalid'
            status_message = "; ".join(errs)
            invalid_count += 1
        elif emp_id:
            plat_lower = (r.get('plataforma') or '').strip().lower()
            usr_lower = (r.get('usuario_login') or '').strip().lower()
            acc_key = (emp_id, plat_lower, usr_lower)
            if acc_key in existing_accounts:
                status = 'update'
                prev_name = existing_accounts[acc_key]['colaborador_nombre']
                status_message = f'Cuenta existente ({prev_name}). Se actualizarán sus datos.'
                update_count += 1
            else:
                ready_count += 1
        else:
            ready_count += 1
            status_message = f"Empresa '{r.get('empresa')}' será creada automáticamente."

        preview_rows.append({
            'index': idx,
            'empresa': r.get('empresa', ''),
            'marca': r.get('marca', ''),
            'plataforma': r.get('plataforma', ''),
            'colaborador_nombre': r.get('colaborador_nombre', ''),
            'colaborador_cargo': r.get('colaborador_cargo', ''),
            'colaborador_email': r.get('colaborador_email', ''),
            'colaborador_telefono': r.get('colaborador_telefono', ''),
            'usuario_login': r.get('usuario_login', ''),
            'password_actual': r.get('password_actual', ''),
            'estado': r.get('estado', 'activo'),
            'notas': r.get('notas', ''),
            'status': status,
            'status_message': status_message,
        })

    return jsonify({
        "success": True,
        "filename": uploaded_file.filename,
        "total_filas": len(preview_rows),
        "ready_count": ready_count,
        "update_count": update_count,
        "invalid_count": invalid_count,
        "rows": preview_rows
    }), 200


@app.route('/api/it-vault/platform-users/import', methods=['POST', 'OPTIONS'])
def api_it_platform_users_import():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    modo = request.form.get('modo', 'crear_o_actualizar')
    created_by = request.form.get('created_by', type=int) or 1

    # También soporta recibir JSON directamente si el frontend confirmó el preview
    payload_rows = []
    if request.is_json:
        data = request.get_json(silent=True) or {}
        payload_rows = data.get('rows', [])
        modo = data.get('modo', modo)
        created_by = data.get('created_by', created_by)
    elif 'file' in request.files:
        uploaded_file = request.files['file']
        if not uploaded_file.filename:
            return jsonify({"error": "Archivo no seleccionado"}), 400
        try:
            payload_rows = parse_accounts_from_filestorage(uploaded_file)
        except Exception as e:
            return jsonify({"error": f"Error al leer el archivo: {str(e)}"}), 400

    if not payload_rows:
        return jsonify({"error": "No se proporcionaron filas para importar"}), 400

    conn = get_db()
    cursor = conn.cursor()
    now_peru = get_peru_now()

    created_count = 0
    updated_count = 0
    skipped_count = 0
    errors = []

    # Cache de empresas, marcas y plataformas
    cursor.execute("SELECT id, LOWER(nombre) FROM it_empresas")
    empresas_map = {r[1]: r[0] for r in cursor.fetchall()}

    cursor.execute("SELECT id, empresa_id, LOWER(nombre) FROM it_marcas")
    marcas_map = {(r[1], r[2]): r[0] for r in cursor.fetchall()}

    cursor.execute("SELECT id, LOWER(nombre) FROM it_plataformas")
    plataformas_map = {r[1]: r[0] for r in cursor.fetchall()}

    for idx, r in enumerate(payload_rows, start=1):
        try:
            emp_name = (r.get('empresa') or '').strip()
            plat_name = (r.get('plataforma') or '').strip()
            colab_name = (r.get('colaborador_nombre') or '').strip()
            usr_login = (r.get('usuario_login') or '').strip()

            if not emp_name or not plat_name or not colab_name or not usr_login:
                errors.append(f"Fila #{idx}: Faltan campos obligatorios (Empresa, Plataforma, Colaborador o Usuario).")
                continue

            # 1. Resolver o Crear Empresa
            emp_lower = emp_name.lower()
            if emp_lower in empresas_map:
                empresa_id = empresas_map[emp_lower]
            else:
                cursor.execute("INSERT INTO it_empresas (nombre, color, created_at) VALUES (?, '#00F0FF', ?)", (emp_name, now_peru))
                empresa_id = cursor.lastrowid
                empresas_map[emp_lower] = empresa_id

            # 2. Resolver o Crear Marca si viene
            marca_name = (r.get('marca') or '').strip()
            marca_id = None
            if marca_name:
                m_key = (empresa_id, marca_name.lower())
                if m_key in marcas_map:
                    marca_id = marcas_map[m_key]
                else:
                    cursor.execute("INSERT INTO it_marcas (empresa_id, nombre, created_at) VALUES (?, ?, ?)", (empresa_id, marca_name, now_peru))
                    marca_id = cursor.lastrowid
                    marcas_map[m_key] = marca_id

            # 3. Registrar plataforma en catálogo si no existe
            plat_lower = plat_name.lower()
            if plat_lower not in plataformas_map:
                cursor.execute("INSERT INTO it_plataformas (nombre, tipo_servicio, color, created_at) VALUES (?, 'General', '#3B82F6', ?)", (plat_name, now_peru))
                plataformas_map[plat_lower] = cursor.lastrowid

            # 4. Resolver contraseña
            raw_password = (r.get('password_actual') or '').strip()
            if not raw_password:
                # Generar contraseña temporal segura si no vino en el Excel
                import secrets
                import string
                chars = string.ascii_letters + string.digits + "!@#$"
                raw_password = "Temp#" + "".join(secrets.choice(chars) for _ in range(8))

            encrypted_pass = encrypt_vault_secret(raw_password)

            cargo = (r.get('colaborador_cargo') or '').strip()
            email = (r.get('colaborador_email') or '').strip()
            telefono = (r.get('colaborador_telefono') or '').strip()
            estado = r.get('estado') or 'activo'
            if estado not in ['activo', 'por_crear', 'suspendido', 'baja']:
                estado = 'activo'
            notas = (r.get('notas') or '').strip()

            # 5. Verificar existencia por (empresa_id, plataforma, usuario_login)
            cursor.execute('''
                SELECT id, password_actual FROM it_platform_users 
                WHERE empresa_id = ? AND LOWER(plataforma) = LOWER(?) AND LOWER(usuario_login) = LOWER(?)
            ''', (empresa_id, plat_name, usr_login))
            existing = cursor.fetchone()

            if existing:
                existing_id = existing[0]
                old_pass = existing[1]
                if modo == 'crear_o_actualizar':
                    cursor.execute('''
                        UPDATE it_platform_users SET
                            marca_id = COALESCE(?, marca_id),
                            colaborador_nombre = ?,
                            colaborador_cargo = ?,
                            colaborador_email = ?,
                            colaborador_telefono = ?,
                            password_anterior = CASE WHEN ? != '' THEN password_actual ELSE password_anterior END,
                            password_actual = CASE WHEN ? != '' THEN ? ELSE password_actual END,
                            estado = ?,
                            notas = ?,
                            updated_by = ?,
                            updated_at = ?
                        WHERE id = ?
                    ''', (marca_id, colab_name, cargo, email, telefono,
                          raw_password, raw_password, encrypted_pass,
                          estado, notas, created_by, now_peru, existing_id))
                    updated_count += 1
                else:
                    skipped_count += 1
            else:
                cursor.execute('''
                    INSERT INTO it_platform_users (
                        empresa_id, marca_id, plataforma, colaborador_nombre, colaborador_cargo,
                        colaborador_email, colaborador_telefono, usuario_login, password_actual,
                        estado, ultimo_reseteo, notas, created_by, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (empresa_id, marca_id, plat_name, colab_name, cargo,
                      email, telefono, usr_login, encrypted_pass,
                      estado, now_peru, notas, created_by, now_peru, now_peru))
                created_count += 1

        except Exception as ex:
            errors.append(f"Fila #{idx}: {str(ex)}")

    conn.commit()
    conn.close()

    total_procesados = created_count + updated_count + skipped_count

    return jsonify({
        "success": True,
        "total_procesados": total_procesados,
        "creados": created_count,
        "actualizados": updated_count,
        "omitidos": skipped_count,
        "errores": errors,
        "mensaje": f"Proceso finalizado: {created_count} creados, {updated_count} actualizados, {skipped_count} omitidos."
    }), 200


# =========================================================================
# RUTAS REST: GESTIÓN DE CATÁLOGOS MAESTROS (PLATAFORMAS, CARGOS, EMPRESAS)
# =========================================================================

@app.route('/api/it-vault/plataformas', methods=['GET', 'POST', 'OPTIONS'])
def api_it_plataformas():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'GET':
        cursor.execute("SELECT * FROM it_plataformas ORDER BY nombre ASC")
        rows = [dict(r) for r in cursor.fetchall()]
        conn.close()
        return jsonify(rows), 200

    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        nombre = (data.get('nombre') or '').strip()
        tipo_servicio = (data.get('tipo_servicio') or 'General').strip()
        color = data.get('color', '#00F0FF')
        if not nombre:
            conn.close()
            return jsonify({"error": "El nombre de la plataforma es obligatorio"}), 400

        try:
            cursor.execute("INSERT INTO it_plataformas (nombre, tipo_servicio, color) VALUES (?, ?, ?)", (nombre, tipo_servicio, color))
            new_id = cursor.lastrowid
            conn.commit()
            conn.close()
            return jsonify({"id": new_id, "nombre": nombre, "tipo_servicio": tipo_servicio, "color": color}), 201
        except sqlite3.IntegrityError:
            conn.close()
            return jsonify({"error": "Esta plataforma ya existe en el catálogo"}), 400


@app.route('/api/it-vault/plataformas/<int:plat_id>', methods=['DELETE', 'OPTIONS'])
def api_it_plataforma_detail(plat_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM it_plataformas WHERE id = ?", (plat_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True}), 200


@app.route('/api/it-vault/cargos', methods=['GET', 'POST', 'OPTIONS'])
def api_it_cargos():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()

    if request.method == 'GET':
        cursor.execute("SELECT * FROM it_cargos ORDER BY nombre ASC")
        rows = [dict(r) for r in cursor.fetchall()]
        conn.close()
        return jsonify(rows), 200

    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        nombre = (data.get('nombre') or '').strip()
        area = (data.get('area') or 'General').strip()
        if not nombre:
            conn.close()
            return jsonify({"error": "El nombre del cargo/área es obligatorio"}), 400

        try:
            cursor.execute("INSERT INTO it_cargos (nombre, area) VALUES (?, ?)", (nombre, area))
            new_id = cursor.lastrowid
            conn.commit()
            conn.close()
            return jsonify({"id": new_id, "nombre": nombre, "area": area}), 201
        except sqlite3.IntegrityError:
            conn.close()
            return jsonify({"error": "Este cargo ya existe en el catálogo"}), 400


@app.route('/api/it-vault/cargos/<int:cargo_id>', methods=['DELETE', 'OPTIONS'])
def api_it_cargo_detail(cargo_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM it_cargos WHERE id = ?", (cargo_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True}), 200


@app.route('/api/it-vault/empresas/<int:empresa_id>', methods=['DELETE', 'OPTIONS'])
def api_it_empresa_delete(empresa_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    # Validar si tiene credenciales o usuarios de plataforma asociados
    cursor.execute("SELECT COUNT(*) FROM it_credentials WHERE empresa_id = ? AND is_active = 1", (empresa_id,))
    if cursor.fetchone()[0] > 0:
        conn.close()
        return jsonify({"error": "No se puede eliminar esta empresa porque contiene credenciales maestras registradas"}), 400

    cursor.execute("SELECT COUNT(*) FROM it_platform_users WHERE empresa_id = ?", (empresa_id,))
    if cursor.fetchone()[0] > 0:
        conn.close()
        return jsonify({"error": "No se puede eliminar esta empresa porque contiene usuarios/colaboradores registrados"}), 400

    cursor.execute("DELETE FROM it_marcas WHERE empresa_id = ?", (empresa_id,))
    cursor.execute("DELETE FROM it_empresas WHERE id = ?", (empresa_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True}), 200


@app.route('/api/it-vault/marcas/<int:marca_id>', methods=['DELETE', 'OPTIONS'])
def api_it_marca_delete(marca_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM it_marcas WHERE id = ?", (marca_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True}), 200


# =========================================================================
# RUTAS REST: MODO BÓVEDA SEGURA (DESBLOQUEO POR PIN / RE-AUTENTICACIÓN)
# =========================================================================

@app.route('/api/it-vault/pin-status', methods=['GET', 'OPTIONS'])
def api_it_vault_pin_status():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    user_id = request.args.get('user_id', type=int)
    if not user_id:
        return jsonify({"error": "user_id es requerido"}), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT security_pin_hash FROM users WHERE id = ?", (user_id,))
    row = cursor.fetchone()
    conn.close()

    has_pin = bool(row and row['security_pin_hash'])
    return jsonify({"has_pin": has_pin}), 200


@app.route('/api/it-vault/set-pin', methods=['POST', 'OPTIONS'])
def api_it_vault_set_pin():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    data = request.get_json(silent=True) or {}
    user_id = data.get('user_id')
    current_password = (data.get('current_password') or '').strip()
    new_pin = (data.get('new_pin') or '').strip()

    if not user_id or not current_password or not new_pin:
        return jsonify({"error": "Contraseña actual y nuevo PIN son obligatorios"}), 400

    if len(new_pin) < 4 or len(new_pin) > 8:
        return jsonify({"error": "El PIN debe tener entre 4 y 8 dígitos o caracteres"}), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT password_hash FROM users WHERE id = ? AND is_active = 1", (user_id,))
    user = cursor.fetchone()
    if not user or not check_password_hash(user['password_hash'], current_password):
        conn.close()
        return jsonify({"error": "La contraseña de inicio de sesión actual es incorrecta"}), 401

    pin_hash = generate_password_hash(new_pin)
    cursor.execute("UPDATE users SET security_pin_hash = ? WHERE id = ?", (pin_hash, user_id))
    
    # Auditoría
    cursor.execute('''
        INSERT INTO it_credential_audit (credential_id, user_id, action, ip_or_agent)
        VALUES (NULL, ?, 'config_security_pin', ?)
    ''', (user_id, request.remote_addr))

    conn.commit()
    conn.close()
    return jsonify({"success": True, "message": "PIN de seguridad configurado exitosamente"}), 200


@app.route('/api/it-vault/verify-unlock', methods=['POST', 'OPTIONS'])
def api_it_vault_verify_unlock():
    if request.method == 'OPTIONS':
        return jsonify({}), 200

    data = request.get_json(silent=True) or {}
    user_id = data.get('user_id')
    password = (data.get('password') or '').strip()
    pin = (data.get('pin') or '').strip()

    if not user_id or (not password and not pin):
        return jsonify({"error": "Debes ingresar tu contraseña o tu PIN de seguridad"}), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id, username, password_hash, security_pin_hash, role FROM users WHERE id = ? AND is_active = 1", (user_id,))
    user = cursor.fetchone()
    if not user:
        conn.close()
        return jsonify({"error": "Usuario no encontrado o inactivo"}), 404

    is_valid = False
    validation_method = None

    if pin:
        if not user['security_pin_hash']:
            conn.close()
            return jsonify({"error": "Aún no has configurado un PIN. Ingresa con tu contraseña de sesión o configúralo."}), 400
        if check_password_hash(user['security_pin_hash'], pin):
            is_valid = True
            validation_method = 'pin'
    elif password:
        if check_password_hash(user['password_hash'], password):
            is_valid = True
            validation_method = 'password'

    if not is_valid:
        # Registrar intento fallido
        cursor.execute('''
            INSERT INTO it_credential_audit (credential_id, user_id, action, ip_or_agent)
            VALUES (NULL, ?, 'failed_vault_unlock', ?)
        ''', (user_id, request.remote_addr))
        conn.commit()
        conn.close()
        return jsonify({"error": "Credencial de desbloqueo incorrecta. Inténtalo de nuevo.", "success": False}), 401

    # Registrar desbloqueo exitoso
    cursor.execute('''
        INSERT INTO it_credential_audit (credential_id, user_id, action, ip_or_agent)
        VALUES (NULL, ?, ?, ?)
    ''', (user_id, f'unlock_vault_{validation_method}', request.remote_addr))
    conn.commit()
    conn.close()

    return jsonify({
        "success": True,
        "expires_in": 300,  # 5 minutos en segundos
        "message": "Bóveda desbloqueada con éxito"
    }), 200

LAST_ROLLOVER_DATE = None
rollover_lock = threading.Lock()

def ensure_daily_rollover(target_date=None):
    global LAST_ROLLOVER_DATE
    today_str = target_date or get_peru_today_str()
    if LAST_ROLLOVER_DATE != today_str:
        with rollover_lock:
            if LAST_ROLLOVER_DATE != today_str:
                print(f"[Scheduler] Verificando y ejecutando rollover para {today_str}...")
                execute_daily_rollover(today_str)
                LAST_ROLLOVER_DATE = today_str

@app.before_request
def check_rollover_on_request():
    if request.path.startswith('/api/') and not request.path.startswith('/api/uploads'):
        ensure_daily_rollover()

def start_midnight_scheduler():
    """Inicia el hilo en background para ejecutar el rollover y catch-up periódico cada 30 segundos."""
    def run_scheduler():
        while True:
            try:
                ensure_daily_rollover()
            except Exception as e:
                print("[Scheduler] Error en ciclo periódico:", e)
            time.sleep(30)

    t = threading.Thread(target=run_scheduler, daemon=True)
    t.start()

# Inicializar base de datos y scheduler tanto en modo directo como con Gunicorn
try:
    init_db()
    cleanup_synthetic_historical_data()
    fix_mock_data_jayala()
    start_midnight_scheduler()
except Exception as e:

    print("[Startup] Error inicializando DB o scheduler:", e)

if __name__ == '__main__':
    print("Iniciando servidor Flask de Bitácoras con soporte de Usuarios, Equipos, Dashboard y Rollover 00:00 en http://localhost:5000...")
    app.run(host='0.0.0.0', port=5000, debug=False)
